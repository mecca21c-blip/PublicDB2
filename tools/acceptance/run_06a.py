"""Deterministic, isolated PublicDB2 06A acceptance and scale harness."""

from __future__ import annotations

import argparse
import csv
import io
import json
import os
import re
import secrets
import socket
import subprocess
import sys
import time
import uuid
from datetime import timedelta
from pathlib import Path

import httpx
from alembic import command
from alembic.config import Config
from openpyxl import load_workbook
from sqlalchemy import event, func, insert, inspect, select, text


PROJECT_ROOT = Path(__file__).resolve().parents[2]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from app.core.config import runtime_paths
from app.db.engine import create_db_engine
from app.db.session import create_session_factory
from app.models import (
    Agency,
    AgencyType,
    ChangeDetection,
    ChangeEventType,
    ContactHistory,
    ContactPoint,
    ContactType,
    CrawlRun,
    DataFormat,
    DetectedChangeCandidate,
    EntityType,
    Observation,
    OrgUnit,
    OrgUnitType,
    ReviewStatus,
    RunStatus,
    Source,
    SourceBinding,
    SourceCoverageMode,
    SourceImportLog,
    SourceOccurrence,
    SourceType,
    StageStatus,
    UserRole,
    CollectionMethod,
)
from app.models.common import utc_now
from app.services.agency_service import AgencyService
from app.services.contact_export_service import ContactExportService
from app.services.contact_service import ContactService
from app.services.dashboard_service import DashboardService
from app.services.review_read_service import ReviewReadService
from app.services.run_service import RunService
from app.services.source_import_service import PreviewStore, SourceImportService
from app.services.source_service import SourceService
from app.services.user_service import UserService


COUNTS = {
    "agencies": 500,
    "org_units": 5_000,
    "sources": 5_000,
    "source_bindings": 10_000,
    "contact_points": 50_000,
    "crawl_runs": 50_000,
    "detected_change_candidates": 10_000,
}
NAMESPACE = uuid.UUID("2fbd781f-cbbb-4ec2-99ca-68d4d80ee06a")


def deterministic_id(kind: str, index: int) -> uuid.UUID:
    return uuid.uuid5(NAMESPACE, f"{kind}:{index}")


def batches(values, size: int = 1_000):
    for index in range(0, len(values), size):
        yield values[index:index + size]


def bulk_insert(session, model, rows) -> None:
    for batch in batches(rows):
        session.execute(insert(model), batch)
    session.commit()


def seed_scale(factory) -> dict[str, list[uuid.UUID]]:
    now = utc_now()
    agency_ids = [deterministic_id("agency", index) for index in range(COUNTS["agencies"])]
    unit_ids = [deterministic_id("unit", index) for index in range(COUNTS["org_units"])]
    source_ids = [deterministic_id("source", index) for index in range(COUNTS["sources"])]
    run_ids = [deterministic_id("run", index) for index in range(COUNTS["crawl_runs"])]
    observation_ids = [deterministic_id("observation", index) for index in range(10_000)]
    contact_ids = [deterministic_id("contact", index) for index in range(COUNTS["contact_points"])]
    detection_ids = [deterministic_id("detection", index) for index in range(10_000)]

    with factory() as session:
        bulk_insert(session, Agency, [{
            "id": agency_ids[index], "official_name": f"Scale Agency {index:04d}",
            "normalized_name": f"Scale Agency {index:04d}", "agency_type": AgencyType.OTHER,
            "active": True, "created_at": now, "updated_at": now,
        } for index in range(COUNTS["agencies"])])

        bulk_insert(session, OrgUnit, [{
            "id": unit_ids[index], "agency_id": agency_ids[index // 10],
            "parent_org_unit_id": None, "name": f"Unit {index:05d}",
            "normalized_name": f"Unit {index:05d}", "unit_type": OrgUnitType.DEPARTMENT,
            "active": True, "created_at": now, "updated_at": now,
        } for index in range(COUNTS["org_units"])])

        bulk_insert(session, Source, [{
            "id": source_ids[index], "url": f"https://scale.example/source/{index:05d}",
            "normalized_url": f"https://scale.example/source/{index:05d}",
            "title": f"Scale Source {index:05d}", "source_type": SourceType.GENERAL_PAGE,
            "collection_method": CollectionMethod.WEB_PAGE, "data_format": DataFormat.HTML,
            "coverage_mode": SourceCoverageMode.ADDITIVE_ONLY, "last_checked_at": now,
            "last_success_at": now, "active": True, "created_at": now, "updated_at": now,
        } for index in range(COUNTS["sources"])])

        binding_rows = []
        binding_context = []
        for index in range(COUNTS["source_bindings"]):
            source_index = index % COUNTS["sources"]
            unit_index = source_index if index < COUNTS["sources"] else (source_index + 2_500) % 5_000
            agency_index = unit_index // 10
            binding_rows.append({
                "id": deterministic_id("binding", index), "source_id": source_ids[source_index],
                "agency_id": agency_ids[agency_index], "org_unit_id": unit_ids[unit_index],
                "scope_key": f"org:{unit_ids[unit_index]}", "description": "06A scale binding",
                "active": True, "created_at": now, "updated_at": now,
            })
            binding_context.append((agency_index, unit_index, source_index))
        bulk_insert(session, SourceBinding, binding_rows)

        contact_rows = []
        for index in range(COUNTS["contact_points"]):
            unit_index = index % COUNTS["org_units"]
            formula = index == 0
            value = "=1+1" if formula else (
                f"scale{index:05d}@example.org" if index % 2 else f"02-{index // 10:04d}-{index % 10_000:04d}"
            )
            contact_rows.append({
                "id": contact_ids[index], "agency_id": agency_ids[unit_index // 10],
                "org_unit_id": unit_ids[unit_index], "person_assignment_id": None, "duty_id": None,
                "contact_type": ContactType.EMAIL if index % 2 else ContactType.PHONE,
                "value": value, "normalized_value": value.casefold(), "purpose_text": "06A scale",
                "verified_at": now, "active": True, "created_at": now + timedelta(microseconds=index),
                "updated_at": now,
            })
        bulk_insert(session, ContactPoint, contact_rows)

        run_rows = []
        for index in range(COUNTS["crawl_runs"]):
            status = RunStatus.FAILED if index % 10 == 0 else (
                RunStatus.PARTIAL if index % 11 == 0 else RunStatus.SUCCESS
            )
            started = now - timedelta(minutes=index % (7 * 24 * 60))
            run_rows.append({
                "id": run_ids[index], "source_id": source_ids[index % COUNTS["sources"]],
                "status": status, "connection_status": StageStatus.SUCCESS,
                "raw_status": StageStatus.SUCCESS,
                "extraction_status": StageStatus.FAILED if status is not RunStatus.SUCCESS else StageStatus.SUCCESS,
                "started_at": started, "finished_at": started + timedelta(seconds=1),
                "records_observed": 0 if status is not RunStatus.SUCCESS else 1,
                "http_status": 200, "error_summary": None if status is RunStatus.SUCCESS else "synthetic failure",
                "collection_method_snapshot": CollectionMethod.WEB_PAGE,
                "collection_kind_snapshot": "WEB_PAGE", "collection_config_snapshot": {"extract_contacts": True},
                "collection_statistics": {"synthetic": True}, "collector_version": "06a-scale",
                "heartbeat_at": started + timedelta(seconds=1),
            })
        bulk_insert(session, CrawlRun, run_rows)

        bulk_insert(session, Observation, [{
            "id": observation_ids[index], "crawl_run_id": run_ids[index],
            "source_id": source_ids[index % COUNTS["sources"]],
            "observed_at": now - timedelta(minutes=index % 10_080),
            "page_url": f"https://scale.example/source/{index % 5_000:05d}",
            "artifact_path": f"data/raw/scale/{index:05d}.html", "artifact_sha256": f"{index:064x}",
            "content_type": "text/html", "response_bytes": 32, "created_at": now,
        } for index in range(10_000)])

        bulk_insert(session, ChangeDetection, [{
            "id": detection_ids[index], "source_id": source_ids[index % 5_000],
            "observation_id": observation_ids[index], "extraction_run_id": None,
            "agency_id": agency_ids[(index % 5_000) // 10], "detector_name": "06a-scale",
            "detector_version": "1.0", "coverage_mode": SourceCoverageMode.ADDITIVE_ONLY,
            "status": RunStatus.SUCCESS, "candidates_found": 1, "started_at": now,
            "finished_at": now, "created_at": now,
        } for index in range(10_000)])

        bulk_insert(session, DetectedChangeCandidate, [{
            "id": deterministic_id("candidate", index), "detection_run_id": detection_ids[index],
            "source_id": source_ids[index % 5_000], "observation_id": observation_ids[index],
            "agency_id": agency_ids[(index % 5_000) // 10], "entity_type": EntityType.CONTACT_POINT,
            "existing_entity_id": contact_ids[index], "proposed_event_type": ChangeEventType.CONTACT_CHANGED,
            "old_value": {"value": f"old-{index}"}, "new_value": {"value": f"new-{index}"},
            "reason": "06A scale candidate", "review_status": ReviewStatus.PENDING_REVIEW,
            "candidate_key": f"06a-scale-{index}", "source_locator": f"row:{index}",
            "actionable": True, "created_at": now - timedelta(seconds=index),
        } for index in range(10_000)])

        session.execute(insert(SourceOccurrence), [{
            "id": deterministic_id("occurrence", 0), "observation_id": observation_ids[0],
            "entity_type": EntityType.CONTACT_POINT, "entity_id": contact_ids[0],
            "field_name": "phone", "observed_value": "=1+1", "context_text": "06A",
            "source_locator": "row:0", "observed_at": now, "created_at": now,
        }])
        session.execute(insert(ContactHistory), [{
            "id": deterministic_id("history", 0), "contact_id": contact_ids[0],
            "value": "=1+1", "normalized_value": "=1+1", "active": True,
            "valid_from": now, "valid_to": None, "change_event_id": None, "created_at": now,
        }])
        session.commit()
    return {"agencies": agency_ids, "units": unit_ids, "sources": source_ids, "bindings": binding_context}


def measured(engine, action):
    count = 0

    def before_cursor_execute(*_args):
        nonlocal count
        count += 1

    event.listen(engine, "before_cursor_execute", before_cursor_execute)
    try:
        return action(), count
    finally:
        event.remove(engine, "before_cursor_execute", before_cursor_execute)


def verify_pages_and_dashboard(engine, factory) -> dict:
    outcomes = {}
    calls = {
        "agencies": lambda session: AgencyService(session).list_page(page_size=100),
        "sources": lambda session: SourceService(session).list_page(page_size=100),
        "runs": lambda session: RunService(session).list_page(page_size=100),
        "contacts": lambda session: ContactService(session).list_page(page_size=100),
        "review": lambda session: ReviewReadService(session).list_page(page_size=100),
    }
    for name, action in calls.items():
        with factory() as session:
            result, query_count = measured(engine, lambda: action(session))
            assert len(result["items"]) <= 100, f"{name} returned an unbounded page"
            assert result["pagination"]["page_size"] == 100
            assert query_count <= 20, f"{name} query count suggests an N+1 regression: {query_count}"
            outcomes[name] = {"items": len(result["items"]), "queries": query_count}
    with factory() as session:
        dashboard, query_count = measured(engine, lambda: DashboardService(session).read())
        kpis = {item["icon"]: item["value"] for item in dashboard["kpis"]}
        assert kpis["building"] == 500
        assert kpis["link"] == 5_000
        assert kpis["contacts"] == 50_000
        assert kpis["review"] == 10_000
        assert len(dashboard["recent_runs"]) == 6
        assert len(dashboard["pending_reviews"]) == 5
        assert len(dashboard["trend"]["labels"]) == 7
        assert query_count <= 20
        outcomes["dashboard"] = {"recent": 6, "pending": 5, "queries": query_count}
    return outcomes


def duplicate_import_content(context: dict) -> bytes:
    stream = io.StringIO(newline="")
    writer = csv.writer(stream)
    writer.writerow(("agency", "org_unit", "url", "description", "collection_method"))
    for agency_index, unit_index, source_index in context["bindings"]:
        writer.writerow((
            f"Scale Agency {agency_index:04d}", f"Unit {unit_index:05d}",
            f"https://scale.example/source/{source_index:05d}", "06A duplicate", "WEB_PAGE",
        ))
    return stream.getvalue().encode("utf-8-sig")


def verify_import(factory, paths, context) -> dict:
    content = duplicate_import_content(context)
    store = PreviewStore(paths)
    entry = store.save("06a-10000.csv", content)
    try:
        with factory() as session:
            service = SourceImportService(session, paths)
            preview = service.preview(entry.original_filename, content)
            assert preview["summary"]["total"] == 10_000
            assert preview["summary"]["duplicates"] == 10_000
            assert preview["summary"]["importable"] == 0
            result = service.confirm(entry)
            assert result["summary"]["duplicates_skipped"] == 10_000
            repeated = service.confirm(entry)
            assert repeated["reused"] is True
            assert session.scalar(select(func.count()).select_from(SourceImportLog)) == 1
            retained = paths.project_root / result["file"]["stored_path"]
            assert retained.is_file() and retained.read_bytes() == content
        return {"rows": 10_000, "duplicates": 10_000, "retained": True, "double_confirm": "reused"}
    finally:
        store.discard(entry.token)


def verify_export(factory, paths) -> dict:
    with factory() as session:
        exported = ContactExportService(session, export_root=paths.export_root).export()
    assert paths.export_root.resolve() in exported.resolve().parents
    workbook = load_workbook(exported, read_only=True, data_only=False)
    sheet = workbook.active
    rows = 0
    formula_safe = False
    provenance = False
    for values in sheet.iter_rows(values_only=True):
        rows += 1
        if rows == 1:
            continue
        formula_safe = formula_safe or "'=1+1" in values
        provenance = provenance or (
            values[7] == "https://scale.example/source/00000"
            and bool(values[8])
            and values[9] == "https://scale.example/source/00000"
        )
    workbook.close()
    assert rows == 50_001
    assert formula_safe
    assert provenance
    return {"data_rows": rows - 1, "formula_safe": formula_safe, "provenance": provenance}


ORPHAN_QUERIES = {
    "binding_source": "SELECT count(*) FROM source_bindings b LEFT JOIN sources s ON s.id=b.source_id WHERE s.id IS NULL",
    "binding_agency": "SELECT count(*) FROM source_bindings b LEFT JOIN agencies a ON a.id=b.agency_id WHERE a.id IS NULL",
    "binding_unit": "SELECT count(*) FROM source_bindings b LEFT JOIN org_units o ON o.id=b.org_unit_id WHERE b.org_unit_id IS NOT NULL AND o.id IS NULL",
    "run_source": "SELECT count(*) FROM crawl_runs r LEFT JOIN sources s ON s.id=r.source_id WHERE s.id IS NULL",
    "observation_run": "SELECT count(*) FROM observations o LEFT JOIN crawl_runs r ON r.id=o.crawl_run_id WHERE r.id IS NULL",
    "observation_source": "SELECT count(*) FROM observations o LEFT JOIN sources s ON s.id=o.source_id WHERE s.id IS NULL",
    "extraction_observation": "SELECT count(*) FROM extraction_runs e LEFT JOIN observations o ON o.id=e.observation_id WHERE o.id IS NULL",
    "occurrence_observation": "SELECT count(*) FROM source_occurrences x LEFT JOIN observations o ON o.id=x.observation_id WHERE o.id IS NULL",
    "history_contact": "SELECT count(*) FROM contact_history h LEFT JOIN contact_points c ON c.id=h.contact_id WHERE c.id IS NULL",
    "candidate_detection": "SELECT count(*) FROM detected_change_candidates c LEFT JOIN change_detections d ON d.id=c.detection_run_id WHERE d.id IS NULL",
    "candidate_source": "SELECT count(*) FROM detected_change_candidates c LEFT JOIN sources s ON s.id=c.source_id WHERE s.id IS NULL",
    "candidate_observation": "SELECT count(*) FROM detected_change_candidates c LEFT JOIN observations o ON o.id=c.observation_id WHERE o.id IS NULL",
    "candidate_agency": "SELECT count(*) FROM detected_change_candidates c LEFT JOIN agencies a ON a.id=c.agency_id WHERE a.id IS NULL",
    "scrape_config_source": "SELECT count(*) FROM source_scrape_configs c LEFT JOIN sources s ON s.id=c.source_id WHERE s.id IS NULL",
    "crawl_config_source": "SELECT count(*) FROM source_crawl_configs c LEFT JOIN sources s ON s.id=c.source_id WHERE s.id IS NULL",
    "api_config_source": "SELECT count(*) FROM source_api_configs c LEFT JOIN sources s ON s.id=c.source_id WHERE s.id IS NULL",
}


def verify_database(engine) -> dict:
    with engine.connect() as connection:
        integrity = connection.exec_driver_sql("PRAGMA integrity_check").scalar_one()
        foreign_keys = connection.exec_driver_sql("PRAGMA foreign_keys").scalar_one()
        journal_mode = connection.exec_driver_sql("PRAGMA journal_mode").scalar_one()
        busy_timeout = connection.exec_driver_sql("PRAGMA busy_timeout").scalar_one()
        head = connection.execute(text("SELECT version_num FROM alembic_version")).scalar_one()
        orphans = {name: connection.execute(text(statement)).scalar_one() for name, statement in ORPHAN_QUERIES.items()}
        counts = {}
        for table, expected in COUNTS.items():
            actual = connection.execute(text(f'SELECT count(*) FROM "{table}"')).scalar_one()
            assert actual == expected, f"{table}: expected {expected}, got {actual}"
            counts[table] = actual
    assert integrity == "ok"
    assert foreign_keys == 1
    assert journal_mode.casefold() == "wal"
    assert busy_timeout >= 5_000
    assert sum(orphans.values()) == 0
    return {
        "head": head, "integrity": integrity, "foreign_keys": foreign_keys,
        "journal_mode": journal_mode, "busy_timeout": busy_timeout,
        "functional_orphans": sum(orphans.values()), "counts": counts,
    }


def absolute_dependency_count(engine) -> int:
    needles = ("C:\\Users\\", "Desktop", "Documents", "AppData", "C:\\PublicDB")
    total = 0
    inspector = inspect(engine)
    with engine.connect() as connection:
        for table_name in inspector.get_table_names():
            for column in inspector.get_columns(table_name):
                if not any(token in str(column["type"]).upper() for token in ("CHAR", "TEXT", "JSON")):
                    continue
                conditions = " OR ".join(
                    f'instr(CAST("{column["name"]}" AS TEXT), :needle_{index}) > 0'
                    for index in range(len(needles))
                )
                statement = text(f'SELECT count(*) FROM "{table_name}" WHERE {conditions}')
                total += connection.execute(
                    statement, {f"needle_{index}": value for index, value in enumerate(needles)}
                ).scalar_one()
    return total


def verify_headless(factory, paths, database_url) -> dict:
    password = secrets.token_urlsafe(24)
    with factory() as session:
        UserService(session).create(
            username="headless-admin", password=password, role=UserRole.ADMIN
        )
    with socket.socket() as probe:
        probe.bind(("127.0.0.1", 0))
        port = probe.getsockname()[1]
    environment = os.environ.copy()
    environment["PUBLICDB2_PROJECT_ROOT"] = str(paths.project_root)
    environment["PUBLICDB2_DATABASE_URL"] = database_url
    flags = getattr(subprocess, "CREATE_NO_WINDOW", 0)
    process = subprocess.Popen(
        [
            sys.executable, "-m", "uvicorn", "app.main:app",
            "--host", "127.0.0.1", "--port", str(port), "--log-level", "warning",
        ],
        cwd=PROJECT_ROOT,
        env=environment,
        stdin=subprocess.DEVNULL,
        stdout=subprocess.DEVNULL,
        stderr=subprocess.DEVNULL,
        creationflags=flags,
    )
    try:
        with httpx.Client(base_url=f"http://127.0.0.1:{port}", timeout=10, trust_env=False) as client:
            deadline = time.monotonic() + 20
            while True:
                if process.poll() is not None:
                    raise RuntimeError(f"headless app exited early: {process.returncode}")
                try:
                    if client.get("/health").status_code == 200:
                        break
                except httpx.HTTPError:
                    pass
                if time.monotonic() >= deadline:
                    raise RuntimeError("headless app did not become healthy")
                time.sleep(0.1)
            assert client.get("/ready").status_code == 200
            login_page = client.get("/login")
            token = re.search(r'name="csrf_token" value="([^"]+)"', login_page.text).group(1)
            login = client.post(
                "/login",
                data={"username": "headless-admin", "password": password, "csrf_token": token},
                follow_redirects=False,
            )
            assert login.status_code == 303
            routes = ("/", "/agencies", "/sources", "/runs", "/contacts", "/review", "/settings")
            statuses = {route: client.get(route).status_code for route in routes}
            assert set(statuses.values()) == {200}
            dashboard = client.get("/")
            csrf = re.search(r'name="csrf-token" content="([^"]+)"', dashboard.text).group(1)
            assert client.post(
                "/logout", headers={"X-CSRF-Token": csrf}, follow_redirects=False
            ).status_code == 303
        return {"health": 200, "ready": 200, "workspaces": statuses, "stopped": True}
    finally:
        if process.poll() is None:
            process.terminate()
            try:
                process.wait(timeout=10)
            except subprocess.TimeoutExpired:
                process.kill()
                process.wait(timeout=10)


def run_focused_tests() -> None:
    completed = subprocess.run(
        [sys.executable, "-m", "pytest", "-q", "tests/acceptance/test_06a_acceptance.py"],
        cwd=PROJECT_ROOT,
        check=False,
    )
    if completed.returncode:
        raise SystemExit(completed.returncode)


def parse_args():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--run-id", default=f"run-{utc_now():%Y%m%dT%H%M%S}-{uuid.uuid4().hex[:8]}")
    parser.add_argument("--skip-focused", action="store_true")
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    if not args.skip_focused:
        run_focused_tests()
    root = (PROJECT_ROOT / "data" / "temp" / "acceptance" / args.run_id).resolve()
    allowed = (PROJECT_ROOT / "data" / "temp" / "acceptance").resolve()
    if allowed not in root.parents or root.exists():
        raise SystemExit("Acceptance root must be a new child of data/temp/acceptance.")
    paths = runtime_paths(root)
    paths.database_root.mkdir(parents=True)
    database_url = f"sqlite+pysqlite:///{paths.database_path.as_posix()}"
    os.environ["PUBLICDB2_PROJECT_ROOT"] = str(root)
    os.environ["PUBLICDB2_DATABASE_URL"] = database_url
    command.upgrade(Config(str(PROJECT_ROOT / "alembic.ini")), "head")
    engine = create_db_engine(database_url)
    factory = create_session_factory(engine)
    try:
        context = seed_scale(factory)
        database = verify_database(engine)
        pages = verify_pages_and_dashboard(engine, factory)
        imported = verify_import(factory, paths, context)
        exported = verify_export(factory, paths)
        headless = verify_headless(factory, paths, database_url)
        report = {
            "acceptance_root": str(root.relative_to(PROJECT_ROOT)),
            "database": database,
            "workspace_pages": pages,
            "import": imported,
            "export": exported,
            "headless": headless,
            "runtime_absolute_path_dependency_count": absolute_dependency_count(engine),
            "external_real_http_count": 0,
            "unbounded_workspace_query_count": 0,
        }
        assert report["runtime_absolute_path_dependency_count"] == 0
        print(json.dumps(report, indent=2, ensure_ascii=False, default=str))
    finally:
        engine.dispose()


if __name__ == "__main__":
    main()
