from __future__ import annotations

import re
from datetime import datetime, timedelta, timezone
from pathlib import Path

import pytest
from alembic import command
from alembic.config import Config
from alembic.autogenerate import compare_metadata
from alembic.migration import MigrationContext
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import inspect

from app.db.base import Base
from app.db.engine import create_db_engine
from app.db.session import create_session_factory
from app.main import create_app
from app.models import (
    Agency, AgencyType, CandidateType, ChangeDetection, ChangeEventType,
    ContactPoint, ContactType, CrawlRun, DetectedChangeCandidate,
    DetectionMethod, EntityType, ExtractedContactCandidate, ExtractionRun,
    ExtractionStatus, Observation, ReviewStatus, RunStatus, Source,
    SourceBinding, SourceCoverageMode, SourceOccurrence, StageStatus, User, UserRole,
)
from app.services.contact_export_service import ContactExportService
from app.services.dashboard_service import DashboardService
from app.services.settings_service import SettingsService, SettingsServiceError
from app.services.user_service import UserService, UserServiceError, verify_password
from tests.support import regression_app


@pytest.fixture
def ops_env(tmp_path, monkeypatch):
    database = tmp_path / 'data' / 'db' / 'publicdb2.sqlite3'
    database.parent.mkdir(parents=True)
    url = f'sqlite:///{database.as_posix()}'
    monkeypatch.setenv('PUBLICDB2_DATABASE_URL', url)
    monkeypatch.setenv('PUBLICDB2_SESSION_SECRET', 'test-secret-' + 'x' * 64)
    config = Config(str(Path(__file__).parents[1] / 'alembic.ini'))
    command.upgrade(config, 'head')
    app = create_app(url, project_root=tmp_path)
    return app, create_session_factory(app.state.engine), tmp_path


def add_user(factory, username, role, password='CorrectHorse12'):
    with factory() as session:
        return UserService(session).create(username=username, password=password, role=role)


def form_csrf(response):
    return re.search(r'name="csrf_token" value="([^"]+)"', response.text).group(1)


def login(client, username, password='CorrectHorse12'):
    token = form_csrf(client.get('/login'))
    return client.post('/login', data={'username': username, 'password': password, 'csrf_token': token}, follow_redirects=False)


def page_csrf(client):
    response = client.get('/')
    return re.search(r'name="csrf-token" content="([^"]+)"', response.text).group(1)


def test_password_hash_generic_failure_and_inactive_login(ops_env):
    app, factory, _ = ops_env
    user = add_user(factory, 'Admin', UserRole.VIEWER)
    with factory() as session:
        stored = session.get(User, user.id)
        assert stored.password_hash != 'CorrectHorse12'
        assert verify_password(stored.password_hash, 'CorrectHorse12')
    with TestClient(app) as client:
        unknown = login(client, 'nobody', 'wrong-password')
        wrong = login(client, 'ADMIN', 'wrong-password')
        assert unknown.status_code == wrong.status_code == 401
        assert '아이디 또는 비밀번호가 올바르지 않습니다.' in unknown.text
        assert '아이디 또는 비밀번호가 올바르지 않습니다.' in wrong.text
    with factory() as session:
        UserService(session).set_active(user.id, False)
    with TestClient(app) as client:
        assert login(client, 'admin').status_code == 401


def test_auth_redirect_roles_csrf_logout_and_safe_payloads(ops_env):
    app, factory, _ = ops_env
    add_user(factory, 'viewer', UserRole.VIEWER)
    add_user(factory, 'operator', UserRole.OPERATOR)
    add_user(factory, 'admin', UserRole.ADMIN)
    with TestClient(app) as client:
        assert client.get('/', follow_redirects=False).status_code == 303
        assert client.get('/api/agencies').status_code == 401
        response = login(client, 'viewer')
        assert response.status_code == 303 and 'httponly' in response.headers['set-cookie'].casefold()
        assert client.get('/agencies').status_code == 200
        token = page_csrf(client)
        assert client.post('/api/agencies', json={'official_name': 'A', 'agency_type': 'OTHER'}).status_code == 403
        assert client.post('/api/agencies', json={'official_name': 'A', 'agency_type': 'OTHER'}, headers={'X-CSRF-Token': token}).status_code == 403
        assert client.post('/api/contacts/export', headers={'X-CSRF-Token': token}).status_code == 403
        assert client.post('/logout', headers={'X-CSRF-Token': token}, follow_redirects=False).status_code == 303
        assert client.get('/api/agencies').status_code == 401
    with TestClient(app) as client:
        login(client, 'operator'); token = page_csrf(client)
        assert client.post('/api/agencies', json={'official_name': 'Operator Agency', 'agency_type': 'OTHER'}, headers={'X-CSRF-Token': token}).status_code == 201
        assert client.put('/api/settings', json={'http_timeout_seconds': 10, 'max_response_bytes': 100000, 'user_agent': 'x'}, headers={'X-CSRF-Token': token}).status_code == 403
        assert client.post('/api/users', json={'username': 'x', 'password': 'long-password-12', 'role': 'VIEWER'}, headers={'X-CSRF-Token': token}).status_code == 403
        assert client.post('/api/contacts/export', headers={'X-CSRF-Token': token}).status_code == 200
    with TestClient(app) as client:
        login(client, 'admin'); token = page_csrf(client)
        saved = client.put('/api/settings', json={'http_timeout_seconds': 11, 'max_response_bytes': 200000, 'user_agent': 'PublicDB2-test'}, headers={'X-CSRF-Token': token})
        created = client.post('/api/users', json={'username': 'new-user', 'display_name': 'New', 'password': 'long-password-12', 'role': 'VIEWER'}, headers={'X-CSRF-Token': token})
        assert saved.status_code == created.status_code == 200 or created.status_code == 201
        assert 'password' not in created.text.casefold() and 'hash' not in created.text.casefold()


def test_last_admin_and_settings_validation_and_live_fetcher(ops_env):
    app, factory, _ = ops_env
    admin = add_user(factory, 'only-admin', UserRole.ADMIN)
    with factory() as session:
        service = UserService(session)
        with pytest.raises(UserServiceError): service.set_active(admin.id, False)
        with pytest.raises(UserServiceError): service.set_role(admin.id, UserRole.OPERATOR)
        settings = SettingsService(session)
        defaults = settings.snapshot()
        assert defaults.http_timeout_seconds > 0 and defaults.max_response_bytes <= 50 * 1024 * 1024
        for values in ((0, 100000, 'agent'), (10, 100, 'agent'), (10, 100000, '')):
            with pytest.raises(SettingsServiceError):
                settings.update(http_timeout_seconds=values[0], max_response_bytes=values[1], user_agent=values[2])
        settings.update(http_timeout_seconds=7, max_response_bytes=123456, user_agent='Configured-Agent')
        service = app.state.collection_service_factory(session)
        assert service.fetcher.timeout_seconds == 7
        assert service.fetcher.max_response_bytes == 123456
        assert service.fetcher.user_agent == 'Configured-Agent'


def test_admin_settings_ui_paths_and_secure_cookie(ops_env, monkeypatch):
    app, factory, root = ops_env
    add_user(factory, 'admin-ui', UserRole.ADMIN)
    monkeypatch.setenv('PUBLICDB2_SECURE_COOKIE', 'true')
    with TestClient(app, base_url='https://testserver') as client:
        response = login(client, 'admin-ui')
        assert 'secure' in response.headers['set-cookie'].casefold()
        page = client.get('/settings')
    assert page.status_code == 200
    for marker in ('사용자 관리', '초기 비밀번호', '역할 변경', '비밀번호 재설정', str(root / 'data' / 'exports')):
        assert marker in page.text
    assert 'password_hash' not in page.text and '$argon2' not in page.text
    assert 'name="project_root"' not in page.text and 'name="export_root"' not in page.text


def test_health_ready_hosts_headers_cors_logs_and_secret_location(ops_env, monkeypatch, tmp_path):
    app, factory, root = ops_env
    add_user(factory, 'health-admin', UserRole.ADMIN)
    with TestClient(app) as client:
        health = client.get('/health')
        ready = client.get('/ready')
        rejected = client.get('/health', headers={'host': 'evil.example'})
    assert health.status_code == 200 and ready.status_code == 200
    assert rejected.status_code == 400
    assert 'sqlite' not in health.text.casefold() and str(root) not in health.text
    assert health.headers['x-content-type-options'] == 'nosniff'
    assert health.headers['x-frame-options'] == 'DENY'
    assert "frame-ancestors 'none'" in health.headers['content-security-policy']
    assert health.headers.get('access-control-allow-origin') != '*'
    assert (root / 'logs' / 'publicdb2.log').is_file()

    bad_root = tmp_path / 'not-ready'
    bad = create_app(f"sqlite:///{(bad_root / 'db.sqlite3').as_posix()}", project_root=bad_root)
    with TestClient(bad) as client:
        assert client.get('/ready').status_code == 503

    monkeypatch.delenv('PUBLICDB2_SESSION_SECRET')
    secret_root = tmp_path / 'secret-owned'
    secret_app = create_app(ops_env[0].state.engine.url.render_as_string(hide_password=False), project_root=secret_root)
    with TestClient(secret_app) as client:
        client.get('/login')
    secret_file = secret_root / 'config' / 'security.json'
    assert secret_file.is_file() and secret_file.stat().st_size > 32


def test_login_csrf_rejects_invalid_token(ops_env):
    app, factory, _ = ops_env
    add_user(factory, 'csrf-admin', UserRole.ADMIN)
    with TestClient(app) as client:
        client.get('/login')
        response = client.post('/login', data={'username': 'csrf-admin', 'password': 'CorrectHorse12', 'csrf_token': 'invalid'})
    assert response.status_code == 403


def test_migration_from_04b_downgrade_reupgrade_and_metadata(tmp_path, monkeypatch):
    database = tmp_path / 'migration.sqlite3'
    url = f'sqlite:///{database.as_posix()}'
    monkeypatch.setenv('PUBLICDB2_DATABASE_URL', url)
    config = Config(str(Path(__file__).parents[1] / 'alembic.ini'))
    command.upgrade(config, '9b2d04b04b01')
    engine = create_db_engine(url)
    assert 'users' not in inspect(engine).get_table_names()
    engine.dispose()
    command.upgrade(config, 'head')
    engine = create_db_engine(url)
    assert {'users', 'operational_settings'} <= set(inspect(engine).get_table_names())
    with engine.connect() as connection:
        assert compare_metadata(MigrationContext.configure(connection), Base.metadata) == []
    engine.dispose()
    command.downgrade(config, '9b2d04b04b01')
    engine = create_db_engine(url)
    assert 'users' not in inspect(engine).get_table_names()
    engine.dispose()
    command.upgrade(config, 'head')
    engine = create_db_engine(url)
    assert {'users', 'operational_settings'} <= set(inspect(engine).get_table_names())
    engine.dispose()


def populate_dashboard(session):
    now = datetime.now(timezone.utc)
    agency = Agency(official_name='Live Agency', normalized_name='live agency', agency_type=AgencyType.OTHER, active=True)
    inactive = Agency(official_name='Inactive', normalized_name='inactive', agency_type=AgencyType.OTHER, active=False)
    session.add_all((agency, inactive)); session.flush()
    first = Source(url='https://example.org/one', normalized_url='https://example.org/one', title='One', active=True)
    second = Source(url='https://example.org/two', normalized_url='https://example.org/two', title='Two', active=True)
    session.add_all((first, second)); session.flush()
    session.add_all((
        SourceBinding(source_id=first.id, agency_id=agency.id, scope_key=f'agency:{agency.id}', active=True),
        SourceBinding(source_id=second.id, agency_id=agency.id, scope_key=f'agency:{agency.id}', active=True),
    ))
    contact = ContactPoint(agency_id=agency.id, contact_type=ContactType.PHONE, value='02-0000-0000', normalized_value='0200000000', active=True, verified_at=now)
    session.add(contact)
    runs = []
    for index in range(8):
        runs.append(CrawlRun(source_id=first.id, status=RunStatus.SUCCESS, connection_status=StageStatus.SUCCESS, raw_status=StageStatus.SUCCESS, extraction_status=StageStatus.SUCCESS, started_at=now-timedelta(days=index), finished_at=now-timedelta(days=index)+timedelta(seconds=1), records_observed=index))
    runs.append(CrawlRun(source_id=second.id, status=RunStatus.PARTIAL, connection_status=StageStatus.SUCCESS, raw_status=StageStatus.SUCCESS, extraction_status=StageStatus.FAILED, started_at=now-timedelta(hours=1), finished_at=now, records_observed=0))
    session.add_all(runs); session.flush()
    observation = Observation(crawl_run_id=runs[0].id, source_id=first.id, observed_at=now, page_url=first.url)
    session.add(observation); session.flush()
    detection = ChangeDetection(source_id=first.id, observation_id=observation.id, agency_id=agency.id, coverage_mode=SourceCoverageMode.ADDITIVE_ONLY, status=RunStatus.SUCCESS, candidates_found=2, started_at=now, finished_at=now)
    session.add(detection); session.flush()
    for index, state in enumerate((ReviewStatus.PENDING_REVIEW, ReviewStatus.DEFERRED, ReviewStatus.APPROVED)):
        session.add(DetectedChangeCandidate(detection_run_id=detection.id, source_id=first.id, observation_id=observation.id, agency_id=agency.id, entity_type=EntityType.SOURCE, existing_entity_id=first.id, proposed_event_type=ChangeEventType.OTHER, old_value='old', new_value=f'new-{index}', reason='changed', review_status=state, candidate_key=f'candidate-{index}', actionable=True, created_at=now-timedelta(minutes=index)))
    session.commit()
    return agency, first, observation, contact


def test_dashboard_live_semantics_recent_trend_empty_and_no_fixture(ops_env):
    app, factory, _ = ops_env
    add_user(factory, 'dashboard-admin', UserRole.ADMIN)
    with factory() as session:
        populate_dashboard(session)
        dashboard = DashboardService(session).read()
    values = {item['label']: item['value'] for item in dashboard['kpis']}
    assert values == {'등록 기관': 1, '수집 소스': 2, '연락처 DB': 1, '검토 대기': 2, '수집 오류': 1}
    assert len(dashboard['recent_runs']) == 6
    assert len(dashboard['trend']['labels']) == 7
    assert dashboard['trend']['success_total'] == 7 and dashboard['trend']['error_total'] == 1
    assert len(dashboard['pending_reviews']) == 2
    with TestClient(app) as client:
        login(client, 'dashboard-admin')
        page = client.get('/')
    assert page.status_code == 200 and '샘플 데이터' not in page.text and 'Live Agency' in page.text

    empty_root = ops_env[2] / 'empty'
    empty_db = empty_root / 'empty.sqlite3'; empty_db.parent.mkdir(parents=True)
    engine = create_db_engine(f'sqlite:///{empty_db.as_posix()}')
    from app.db.base import Base
    Base.metadata.create_all(engine)
    with create_session_factory(engine)() as session:
        empty = DashboardService(session).read()
    engine.dispose()
    assert all(item['value'] == 0 for item in empty['kpis'])
    assert empty['recent_runs'] == () and empty['pending_reviews'] == ()


def test_dashboard_db_error_never_falls_back_to_fixture(tmp_path):
    url = f"sqlite:///{(tmp_path/'missing.sqlite3').as_posix()}"
    with TestClient(regression_app(url, project_root=tmp_path)) as client:
        page = client.get('/')
    assert page.status_code == 200
    assert '대시보드 데이터를 불러오지 못했습니다.' in page.text
    assert '샘플 데이터' not in page.text


def test_confirmed_export_filters_provenance_formula_and_bounds(ops_env):
    app, factory, root = ops_env
    add_user(factory, 'export-operator', UserRole.OPERATOR)
    with factory() as session:
        agency, source, observation, contact = populate_dashboard(session)
        contact.value = '=HYPERLINK("bad")'
        contact.normalized_value = '=hyperlinkbad'
        session.add(SourceOccurrence(observation_id=observation.id, entity_type=EntityType.CONTACT_POINT, entity_id=contact.id, field_name='value', observed_value=contact.value, observed_at=observation.observed_at))
        extraction = ExtractionRun(observation_id=observation.id, extractor_name='html_contact', extractor_version='test', status=ExtractionStatus.SUCCESS, started_at=observation.observed_at, finished_at=observation.observed_at, candidates_found=1)
        session.add(extraction); session.flush()
        session.add(ExtractedContactCandidate(extraction_run_id=extraction.id, observation_id=observation.id, candidate_type=CandidateType.EMAIL, raw_value='discovery-only@example.org', normalized_value='discovery-only@example.org', detection_method=DetectionMethod.TEXT_PATTERN))
        other = Agency(official_name='Other Agency', normalized_name='other agency', agency_type=AgencyType.OTHER, active=True)
        session.add(other); session.flush()
        session.add(ContactPoint(agency_id=other.id, contact_type=ContactType.EMAIL, value='other@example.org', normalized_value='other@example.org', active=True))
        session.commit()
        path = ContactExportService(session, export_root=root/'data'/'exports').export(search='Live Agency', agency_id=agency.id, contact_type=ContactType.PHONE)
    assert (root/'data'/'exports').resolve() in path.parents
    assert path.suffix == '.xlsx' and path.name.startswith('publicdb2_contacts_')
    sheet = load_workbook(path, read_only=True).active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows[0] == ('기관', '부서', '업무', '담당자', '연락처 유형', '연락처 값', '확인일', '공식 출처 URL')
    assert len(rows) == 2 and rows[1][0] == 'Live Agency'
    assert rows[1][5].startswith("'=")
    assert rows[1][7] == 'https://example.org/one'
    assert 'discovery-only@example.org' not in str(rows)

    with TestClient(app) as client:
        login(client, 'export-operator'); token = page_csrf(client)
        response = client.post('/api/contacts/export?path=C:/outside.xlsx&agency_id=' + str(agency.id), headers={'X-CSRF-Token': token})
    assert response.status_code == 200
    assert not (root/'outside.xlsx').exists()
