from __future__ import annotations

import base64
import csv
import io
import json
import subprocess
import uuid
from pathlib import Path
from urllib.parse import urlparse

import pytest
from alembic import command
from alembic.config import Config
from bs4 import BeautifulSoup
from fastapi.testclient import TestClient
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app.core.config import runtime_paths
from app.db.engine import create_db_engine
from app.db.session import create_session_factory
from app.models import (
    Agency, AgencyType, ApiAuthMode, ApiPaginationMode, ApiSourceKind,
    CollectionJob, CollectionJobItem, CollectionMethod, ContactPoint, CrawlRun,
    Duty, OrgUnit, Person, PersonAssignment, Source, SourceBinding, SourceImportLog,
)
from app.services.agency_service import AgencyService
from app.services.collection_job_service import CollectionJobService
from app.services.source_import_service import (
    MAX_IMPORT_ROWS, PREVIEW_PAGE_SIZE, PreviewStore, SourceImportError,
    SourceImportService, parse_import_file,
)
from app.services.source_service import SourceService
from tests.support import regression_app


V2_HEADERS = (
    "기관명", "기관유형", "지역", "부서명", "URL", "수집방식", "소스설명", "자동수집",
)


@pytest.fixture()
def astra_env(tmp_path, monkeypatch):
    database = tmp_path / "db" / "astra.sqlite3"
    database.parent.mkdir(parents=True)
    url = f"sqlite:///{database.as_posix()}"
    monkeypatch.setenv("PUBLICDB2_DATABASE_URL", url)
    command.upgrade(Config("alembic.ini"), "head")
    engine = create_db_engine(url)
    factory = create_session_factory(engine)
    yield tmp_path, url, factory
    engine.dispose()


def csv_v2(rows, headers=V2_HEADERS) -> bytes:
    stream = io.StringIO()
    writer = csv.writer(stream)
    writer.writerow(headers)
    writer.writerows(rows)
    return ("\ufeff" + stream.getvalue()).encode()


def xlsx_v2(rows, headers=V2_HEADERS) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(list(headers))
    for row in rows:
        sheet.append(list(row))
    stream = io.BytesIO()
    workbook.save(stream)
    return stream.getvalue()


def confirm_file(session, root: Path, filename: str, content: bytes, **options):
    store = PreviewStore(runtime_paths(root))
    entry = store.save(filename, content)
    result = SourceImportService(session, runtime_paths(root)).confirm(entry, **options)
    return store, entry, result


def test_v2_xlsx_korean_aliases_agency_region_methods_and_refresh(astra_env):
    root, _url, factory = astra_env
    rows = [
        ["중앙 기관", "중앙부처", "", "", "https://example.org/page", "스크래핑", "페이지", "Y"],
        ["서울 광역", "광역지자체", "서울", "정보과", "https://example.org/crawl", "크롤링", "크롤", "제외"],
        ["공공 기관", "공공기관", "", "", "https://example.org/rss?kind=notice", "RSS", "피드", ""],
        ["위원 기관", "위원회", "SEOUL", "", "https://example.org/atom", "Atom", "Atom", "자동"],
    ]
    with factory() as session:
        service = SourceImportService(session, runtime_paths(root))
        preview = service.preview("astra.xlsx", xlsx_v2(rows))
        assert [row["classification"] for row in preview["rows"]] == ["NEW_AGENCY"] * 4
        assert [row["agency_type"] for row in preview["rows"]] == [
            "CENTRAL_GOVERNMENT", "METROPOLITAN_GOVERNMENT", "PUBLIC_INSTITUTION", "COMMISSION",
        ]
        assert [row["region_code"] for row in preview["rows"]] == [None, "SEOUL", None, "SEOUL"]
        assert [row["scheduled_refresh_enabled"] for row in preview["rows"]] == [True, False, True, True]
        assert [(row["collection_method"], row["api_kind"]) for row in preview["rows"]] == [
            ("WEB_PAGE", None), ("WEB_CRAWL", None), ("API", "RSS"), ("API", "ATOM"),
        ]

        store, entry, result = confirm_file(session, root, "astra.xlsx", xlsx_v2(rows))
        assert result["summary"]["created_agencies"] == 4
        assert result["summary"]["created_org_units"] == 1
        assert result["summary"]["created_sources"] == 4
        rss = session.scalar(select(Source).where(Source.url.contains("rss?kind=notice")))
        atom = session.scalar(select(Source).where(Source.url.endswith("/atom")))
        assert rss.url.endswith("?kind=notice") and rss.api_config.kind is ApiSourceKind.RSS
        assert atom.api_config.kind is ApiSourceKind.ATOM
        for source in (rss, atom):
            assert source.api_config.auth_mode is ApiAuthMode.NONE
            assert source.api_config.pagination_mode is ApiPaginationMode.NONE
            assert source.api_config.discovery_only is True
        assert session.scalar(select(Source).where(Source.url.endswith("/crawl"))).crawl_config is not None
        assert session.scalar(select(Source).where(Source.url.endswith("/page"))).scrape_config is not None
        store.discard(entry.token)


def test_new_agency_and_org_are_uniquely_planned_and_confirmed(astra_env):
    root, _url, factory = astra_env
    rows = [
        ["신규테스트기관", "PUBLIC_INSTITUTION", "", "정보과", f"https://unique.example/{index}", "WEB_PAGE", str(index), "TRUE"]
        for index in range(3)
    ]
    with factory() as session:
        service = SourceImportService(session, runtime_paths(root))
        preview = service.preview("unique.csv", csv_v2(rows))
        assert preview["summary"]["new_agencies"] == 1
        assert preview["summary"]["new_org_units"] == 1
        assert preview["summary"]["new_sources"] == 3
        store, entry, result = confirm_file(session, root, "unique.csv", csv_v2(rows))
        assert result["summary"]["created_agencies"] == 1
        assert result["summary"]["created_org_units"] == 1
        assert result["summary"]["created_sources"] == 3
        assert result["summary"]["created_bindings"] == 3
        assert session.scalar(select(func.count()).select_from(Agency)) == 1
        assert session.scalar(select(func.count()).select_from(OrgUnit)) == 1
        store.discard(entry.token)


def test_existing_agency_metadata_is_preserved_and_conflicts_are_isolated(astra_env):
    root, _url, factory = astra_env
    with factory() as session:
        agency, _ = AgencyService(session).create_agency(
            official_name="서울특별시", agency_type=AgencyType.METROPOLITAN_GOVERNMENT,
            region_code="SEOUL",
        )
        rows = [
            ["서울특별시", "", "", "", "https://meta.example/blank", "WEB_PAGE", "blank", "N"],
            ["서울특별시", "광역자치단체", "서울", "", "https://meta.example/same", "WEB_PAGE", "same", "Y"],
            ["서울특별시", "PUBLIC_INSTITUTION", "BUSAN", "", "https://meta.example/conflict", "WEB_PAGE", "bad", "Y"],
            ["지역 누락 구청", "BASIC_LOCAL_GOVERNMENT", "", "", "https://meta.example/missing", "WEB_PAGE", "bad", "Y"],
        ]
        preview = SourceImportService(session, runtime_paths(root)).preview("metadata.csv", csv_v2(rows))
        assert [row["classification"] for row in preview["rows"]] == ["READY", "READY", "CONFLICT", "INVALID"]
        store, entry, result = confirm_file(session, root, "metadata.csv", csv_v2(rows))
        assert result["summary"]["created_sources"] == 2
        session.expire_all()
        existing = session.get(Agency, uuid.UUID(agency["id"]))
        assert existing.agency_type is AgencyType.METROPOLITAN_GOVERNMENT
        assert existing.region_code == "SEOUL"
        store.discard(entry.token)


def test_openapi_method_and_api_kind_conflicts_are_explicit(astra_env):
    root, _url, factory = astra_env
    rows = [
        ["기관", "OTHER", "", "", "https://kind.example/open", "OpenAPI", "", "Y"],
        ["기관", "OTHER", "", "", "https://kind.example/shared", "RSS", "", "Y"],
        ["기관", "OTHER", "", "", "https://kind.example/shared", "ATOM", "", "Y"],
        ["기관", "OTHER", "", "", "https://kind.example/method", "WEB_PAGE", "", "Y"],
        ["기관", "OTHER", "", "", "https://kind.example/method", "WEB_CRAWL", "", "Y"],
    ]
    with factory() as session:
        preview = SourceImportService(session, runtime_paths(root)).preview("conflicts.csv", csv_v2(rows))
        assert [row["classification"] for row in preview["rows"]] == [
            "INVALID", "NEW_AGENCY", "CONFLICT", "NEW_AGENCY", "CONFLICT",
        ]
        assert "전용 Source 설정" in preview["rows"][0]["message"]
        assert "API kind" in preview["rows"][2]["message"]


def test_existing_source_new_binding_keeps_scheduled_flag_and_description(astra_env):
    root, _url, factory = astra_env
    with factory() as session:
        agencies = AgencyService(session)
        first, _ = agencies.create_agency(official_name="기존 기관", agency_type=AgencyType.OTHER)
        second, _ = agencies.create_agency(official_name="연결 기관", agency_type=AgencyType.OTHER)
        SourceService(session).register_binding(
            url="https://existing.example/source", agency_id=uuid.UUID(first["id"]),
            scheduled_refresh_enabled=False,
        )
        rows = [["연결 기관", "", "", "", "https://existing.example/source", "WEB_PAGE", "ASTRA 설명", "Y"]]
        store, entry, result = confirm_file(session, root, "binding.csv", csv_v2(rows))
        assert result["summary"]["created_sources"] == 0
        assert result["summary"]["existing_source_new_bindings"] == 1
        source = session.scalar(select(Source).where(Source.normalized_url == "https://existing.example/source"))
        assert source.scheduled_refresh_enabled is False
        binding = session.scalar(select(SourceBinding).where(SourceBinding.agency_id == uuid.UUID(second["id"])))
        assert binding.description == "ASTRA 설명"
        store.discard(entry.token)


def test_preview_is_write_free_filtered_and_bounded_for_one_thousand_rows(astra_env):
    root, _url, factory = astra_env
    rows = [
        [f"기관 {index % 10}", "PUBLIC_INSTITUTION", "", f"부서 {index % 20}",
         f"https://scale.example/{index}", ("WEB_PAGE", "WEB_CRAWL", "RSS", "ATOM")[index % 4], "", "Y"]
        for index in range(800)
    ]
    rows += [list(rows[index]) for index in range(100)]
    rows += [["오류 기관", "OTHER", "", "", f"ftp://scale.example/{index}", "WEB_PAGE", "", "Y"] for index in range(50)]
    for index in range(50):
        conflict = list(rows[index])
        conflict[5] = "WEB_CRAWL" if conflict[5] == "WEB_PAGE" else "WEB_PAGE"
        rows.append(conflict)
    content = csv_v2(rows)
    with factory() as session:
        models = (Agency, OrgUnit, Source, SourceBinding, CollectionJob, CrawlRun, SourceImportLog)
        before = {model: session.scalar(select(func.count()).select_from(model)) for model in models}
        service = SourceImportService(session, runtime_paths(root))
        preview = service.preview("scale.csv", content)
        assert preview["summary"] == {
            "total": 1000, "importable": 800, "new_agencies": 10,
            "new_org_units": 20, "new_sources": 800,
            "existing_source_new_bindings": 0, "duplicates": 100,
            "conflicts": 50, "invalid": 50,
            "categories": {
                "READY": 0, "NEW_AGENCY": 800, "NEW_ORG_UNIT": 0,
                "EXACT_DUPLICATE": 100, "EXISTING_SOURCE_NEW_BINDING": 0,
                "CONFLICT": 50, "INVALID": 50,
            },
        }
        assert len(preview["rows"]) == PREVIEW_PAGE_SIZE
        assert preview["preview"]["pages"] == 20
        conflicts = service.preview("scale.csv", content, filter_name="conflict")
        assert len(conflicts["rows"]) == 50
        after = {model: session.scalar(select(func.count()).select_from(model)) for model in models}
        assert after == before

        store, entry, result = confirm_file(session, root, "scale.csv", content)
        assert result["summary"] == {
            "input_rows": 1000, "created_agencies": 10, "created_org_units": 20,
            "created_sources": 800, "created_bindings": 800,
            "existing_source_new_bindings": 0, "duplicates_skipped": 100,
            "invalid_rows": 50, "conflicts": 50, "unexpected_failures": 0,
        }
        assert session.scalar(select(func.count()).select_from(Agency)) == 10
        assert session.scalar(select(func.count()).select_from(OrgUnit)) == 20
        assert session.scalar(select(func.count()).select_from(Source)) == 800
        assert session.scalar(select(func.count()).select_from(SourceBinding)) == 800
        for model in (Duty, Person, PersonAssignment, ContactPoint):
            assert session.scalar(select(func.count()).select_from(model)) == 0
        store.discard(entry.token)


def test_ten_thousand_row_parser_limit_and_no_dom_projection(astra_env):
    root, _url, factory = astra_env
    rows = [["기관", "OTHER", "", "", f"https://limit.example/{index}", "WEB_PAGE", "", "Y"] for index in range(MAX_IMPORT_ROWS)]
    content = csv_v2(rows)
    assert len(parse_import_file("limit.csv", content)) == MAX_IMPORT_ROWS
    with pytest.raises(SourceImportError, match="최대"):
        parse_import_file("too-many.csv", csv_v2(rows + [["기관", "OTHER", "", "", "https://limit.example/extra", "WEB_PAGE", "", "Y"]]))
    with factory() as session:
        preview = SourceImportService(session, runtime_paths(root)).preview("limit.csv", content)
        assert preview["summary"]["total"] == MAX_IMPORT_ROWS
        assert len(preview["rows"]) == PREVIEW_PAGE_SIZE


def test_collect_after_confirm_queues_one_job_for_unique_new_sources_and_is_idempotent(astra_env):
    root, _url, factory = astra_env
    valid = [
        ["Queue 기관", "PUBLIC_INSTITUTION", "", "", f"https://queue.example/{index}", "WEB_PAGE", "", "Y"]
        for index in range(80)
    ]
    rows = valid + [list(valid[index]) for index in range(10)]
    rows += [["Queue 기관", "PUBLIC_INSTITUTION", "", "", f"ftp://queue.example/{index}", "WEB_PAGE", "", "Y"] for index in range(10)]
    with factory() as session:
        agencies = AgencyService(session)
        owner, _ = agencies.create_agency(official_name="기존 소유 기관", agency_type=AgencyType.OTHER)
        linked, _ = agencies.create_agency(official_name="기존 새 연결 기관", agency_type=AgencyType.OTHER)
        SourceService(session).register_binding(
            url="https://queue.example/existing", agency_id=uuid.UUID(owner["id"]),
        )
        rows.append([
            "기존 새 연결 기관", "OTHER", "", "", "https://queue.example/existing",
            "WEB_PAGE", "기존 Source 새 연결", "Y",
        ])
        content = csv_v2(rows)
        store = PreviewStore(runtime_paths(root))
        entry = store.save("queue.csv", content)
        service = SourceImportService(session, runtime_paths(root))
        first = service.confirm(entry, collect_after_confirm=True)
        second = service.confirm(entry, collect_after_confirm=True)
        assert first["summary"]["created_sources"] == 80
        assert first["summary"]["existing_source_new_bindings"] == 1
        assert first["collection_requested"] is True
        assert first["collection_source_count"] == 80
        assert first["collection_queue_status"] == "QUEUED"
        assert second["reused"] is True
        assert second["collection_job_id"] == first["collection_job_id"]
        assert session.scalar(select(func.count()).select_from(CollectionJob)) == 1
        assert session.scalar(select(func.count()).select_from(CollectionJobItem)) == 80
        job = session.scalar(select(CollectionJob))
        assert job.trigger_context["origin"] == "SOURCE_IMPORT"
        assert job.trigger_context["new_source_count"] == 80
        assert len(job.trigger_context) < 10
        store.discard(entry.token)


def test_collect_false_creates_no_job_and_queue_failure_does_not_rollback(astra_env, monkeypatch):
    root, _url, factory = astra_env
    content = csv_v2([["기관", "OTHER", "", "", "https://failure.example/source", "WEB_PAGE", "", "Y"]])
    with factory() as session:
        store, entry, result = confirm_file(session, root, "off.csv", content, collect_after_confirm=False)
        assert result["collection_queue_status"] == "NOT_REQUESTED"
        assert session.scalar(select(func.count()).select_from(CollectionJob)) == 0
        store.discard(entry.token)

    second_content = csv_v2([["기관 2", "OTHER", "", "", "https://failure.example/second", "WEB_PAGE", "", "Y"]])
    monkeypatch.setattr(
        CollectionJobService, "create_from_sources",
        lambda *_args, **_kwargs: (_ for _ in ()).throw(RuntimeError("queue unavailable")),
    )
    with factory() as session:
        store, entry, result = confirm_file(session, root, "failure.csv", second_content, collect_after_confirm=True)
        assert result["collection_queue_status"] == "FAILED_TO_QUEUE"
        assert result["summary"]["created_sources"] == 1
        assert session.scalar(select(func.count()).select_from(Source)) == 2
        assert session.scalar(select(func.count()).select_from(SourceBinding)) == 2
        assert session.scalar(select(func.count()).select_from(CollectionJob)) == 0
        assert session.scalar(select(func.count()).select_from(SourceImportLog)) == 2
        store.discard(entry.token)


def test_template_routes_and_bounded_preview_api(astra_env):
    root, url, _factory = astra_env
    app = regression_app(url, project_root=root)
    content = xlsx_v2([
        ["API 기관", "PUBLIC_INSTITUTION", "", "", f"https://api.example/{index}", "WEB_PAGE", "", "Y"]
        for index in range(120)
    ])
    with TestClient(app) as client:
        xlsx = client.get("/api/source-bindings/imports/template.xlsx")
        csv_template = client.get("/api/source-bindings/imports/template.csv")
        preview = client.post("/api/source-bindings/imports/preview", files={
            "file": ("astra.xlsx", content, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        })
        token = preview.json()["token"]
        page_two = client.get(f"/api/source-bindings/imports/{token}/preview?filter=all&page=2&page_size=50")
        client.delete(f"/api/source-bindings/imports/{token}")
    assert xlsx.status_code == 200
    workbook = load_workbook(io.BytesIO(xlsx.content), read_only=True)
    assert workbook.sheetnames == ["수집 소스", "입력 안내"]
    assert tuple(cell.value for cell in next(workbook["수집 소스"].iter_rows(max_row=1))) == V2_HEADERS
    assert csv_template.content.startswith(b"\xef\xbb\xbf")
    assert preview.status_code == 200 and len(preview.json()["rows"]) == 50
    assert page_two.status_code == 200 and page_two.json()["preview"]["page"] == 2
    assert len(page_two.json()["rows"]) == 50


def test_import_modal_production_dom_headless_path(astra_env):
    root, url, _factory = astra_env
    page = TestClient(regression_app(url, project_root=root)).get("/sources")
    browser = next((candidate for candidate in (
        Path(r"C:\Program Files\Google\Chrome\Application\chrome.exe"),
        Path(r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"),
    ) if candidate.exists()), None)
    if browser is None:
        pytest.skip("Installed Chromium is required for the Astra import browser contract")

    soup = BeautifulSoup(page.text, "html.parser")
    for link in list(soup.select('link[rel="stylesheet"]')):
        static_path = urlparse(link["href"]).path
        if static_path.startswith("/static/"):
            style = soup.new_tag("style")
            style.string = (Path("app/web/static") / static_path.removeprefix("/static/")).read_text(encoding="utf-8")
            link.replace_with(style)
    app_js = Path("app/web/static/js/app.js").read_text(encoding="utf-8")
    for script in list(soup.select('script[src$="/static/js/app.js"]')):
        inline = soup.new_tag("script")
        inline.string = app_js
        script.replace_with(inline)
    encoded = base64.b64encode(xlsx_v2([["기관", "OTHER", "", "", "https://browser.example/source", "WEB_PAGE", "", "Y"]])).decode()
    harness = soup.new_tag("script")
    harness.string = f"""
      setTimeout(async () => {{
        const delay = (ms) => new Promise((resolve) => setTimeout(resolve, ms));
        const response = (payload) => ({{ok: true, status: 200, json: async () => payload}});
        const makeRows = (start, count, classification = 'NEW_AGENCY') => Array.from({{length: count}}, (_, index) => ({{
          row_number: start + index, raw_agency: 'ASTRA 기관', resolved_agency: 'ASTRA 기관',
          agency_type: 'PUBLIC_INSTITUTION', region_code: null, resolved_org_unit: null,
          raw_url: 'https://browser.example/' + (start + index), normalized_url: 'https://browser.example/' + (start + index),
          collection_method: 'WEB_PAGE', api_kind: null, scheduled_refresh_enabled: true,
          classification, message: classification === 'INVALID' ? 'URL 오류' : '기관을 새로 만듭니다.', importable: classification !== 'INVALID',
        }}));
        const summary = {{total: 120, importable: 117, new_agencies: 1, new_org_units: 0, new_sources: 117,
          existing_source_new_bindings: 0, duplicates: 0, conflicts: 0, invalid: 3}};
        window.fetch = async (target, options = {{}}) => {{
          const value = String(target);
          if (value.endsWith('/imports/preview')) return response({{token: 'preview-token', filename: 'astra.xlsx', rows: makeRows(2, 50), summary,
            preview: {{filter: 'all', page: 1, page_size: 50, total: 120, pages: 3, has_previous: false, has_next: true}}}});
          if (value.includes('/preview?')) {{
            const query = new URL(value, 'https://local.test').searchParams;
            if (query.get('filter') === 'invalid') return response({{rows: makeRows(119, 3, 'INVALID'), summary,
              preview: {{filter: 'invalid', page: 1, page_size: 50, total: 3, pages: 1, has_previous: false, has_next: false}}}});
            return response({{rows: makeRows(52, 50), summary,
              preview: {{filter: 'all', page: 2, page_size: 50, total: 120, pages: 3, has_previous: true, has_next: true}}}});
          }}
          if (value.includes('/confirm?')) return response({{summary: {{input_rows: 120, created_agencies: 1, created_org_units: 0,
            created_sources: 117, created_bindings: 117, existing_source_new_bindings: 0, duplicates_skipped: 0,
            conflicts: 0, invalid_rows: 3, unexpected_failures: 0}}, collection_requested: true,
            collection_source_count: 117, collection_queue_status: 'QUEUED', collection_job_id: 'job-1'}});
          return response({{job: null, jobs: []}});
        }};
        try {{
          document.querySelector('[data-modal-open="excel-import"]').click();
          const bytes = Uint8Array.from(atob('{encoded}'), (char) => char.charCodeAt(0));
          const transfer = new DataTransfer();
          transfer.items.add(new File([bytes], 'astra.xlsx', {{type: 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'}}));
          const input = document.querySelector('[data-import-preview] input[type="file"]');
          input.files = transfer.files;
          document.querySelector('[data-import-preview] button[type="submit"]').click();
          await delay(120);
          const initialRows = document.querySelectorAll('[data-import-rows] tr').length;
          document.querySelector('[data-import-next]').click();
          await delay(80);
          const secondPage = document.querySelector('[data-import-page-status]').textContent;
          const filter = document.querySelector('[data-import-filter]');
          filter.value = 'invalid'; filter.dispatchEvent(new Event('change', {{bubbles: true}}));
          await delay(80);
          const filteredRows = document.querySelectorAll('[data-import-rows] tr').length;
          document.querySelector('[data-import-collect-after-confirm]').click();
          document.querySelector('[data-import-confirm]').click();
          await delay(100);
          document.body.dataset.astraImportContract = JSON.stringify({{
            initialRows, secondPage, filteredRows,
            completed: !document.querySelector('[data-import-complete]').hidden,
            queued: document.querySelector('[data-import-collection-result]').textContent,
            historyVisible: !document.querySelector('[data-import-history]').hidden,
            objectText: document.querySelector('[data-import-workflow]').innerText.includes('[object Object]'),
          }});
        }} catch (error) {{ document.body.dataset.astraImportError = error.stack || String(error); }}
      }}, 50);
    """
    soup.body.append(harness)
    html_path = root / "astra-import-browser.html"
    html_path.write_text(str(soup), encoding="utf-8")
    completed = subprocess.run([
        str(browser), "--headless=new", "--disable-gpu", "--no-first-run", "--no-default-browser-check",
        f"--user-data-dir={root / 'chrome-profile'}", "--virtual-time-budget=4000", "--dump-dom", html_path.as_uri(),
    ], check=True, capture_output=True, text=True, encoding="utf-8", timeout=30)
    rendered = BeautifulSoup(completed.stdout, "html.parser")
    assert "data-astra-import-error" not in rendered.body.attrs, rendered.body.get("data-astra-import-error")
    contract = json.loads(rendered.body["data-astra-import-contract"])
    assert contract == {
        "initialRows": 50, "secondPage": "120건 중 2 / 3 페이지", "filteredRows": 3,
        "completed": True, "queued": "신규 Source 117건의 수집 작업이 대기 중입니다.",
        "historyVisible": True, "objectText": False,
    }
