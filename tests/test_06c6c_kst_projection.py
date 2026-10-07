from __future__ import annotations

from datetime import datetime, timezone
import uuid

from alembic import command
from alembic.config import Config
from openpyxl import load_workbook

from app.db.engine import create_db_engine
from app.db.session import create_session_factory

from app.core.time_presentation import format_kst_datetime, serialize_utc_datetime, to_kst
from app.models import (
    AgencyType, ChangeDetection, ChangeEventType, CollectionJob,
    CollectionJobStatus, CollectionTriggerType, ContactHistory, ContactPoint,
    ContactType, CrawlRun, DetectedChangeCandidate, EntityType, Observation,
    ReviewStatus, RunStatus, SourceCoverageMode, SourceOccurrence, StageStatus,
)
from app.models.common import UTCDateTime
from app.services.agency_service import AgencyService
from app.services.collection_job_service import CollectionJobService
from app.services.contact_export_service import ContactExportService
from app.services.contact_service import ContactService
from app.services.review_read_service import ReviewReadService
from app.services.run_service import RunService
from app.services.source_service import SourceService


def test_kst_projection_same_date_aware_utc_and_none():
    value = datetime(2026, 10, 7, 8, 22, tzinfo=timezone.utc)

    assert format_kst_datetime(value) == "2026-10-07 17:22"
    assert to_kst(value).utcoffset().total_seconds() == 9 * 60 * 60
    assert format_kst_datetime(None) == "-"


def test_kst_projection_rolls_over_midnight_and_treats_naive_as_utc():
    naive_utc = datetime(2026, 10, 7, 15, 30)

    assert format_kst_datetime(naive_utc) == "2026-10-08 00:30"
    assert serialize_utc_datetime(naive_utc) == "2026-10-07T15:30:00+00:00"


def test_sqlite_result_contract_is_aware_utc_before_kst_projection():
    raw_sqlite_value = datetime(2026, 10, 7, 8, 22)

    returned = UTCDateTime().process_result_value(raw_sqlite_value, None)

    assert returned.tzinfo is timezone.utc
    assert format_kst_datetime(returned) == "2026-10-07 17:22"


def test_collection_job_keeps_explicit_utc_api_and_adds_kst_display():
    created = datetime(2026, 10, 7, 15, 30, tzinfo=timezone.utc)
    job = CollectionJob(
        trigger_type=CollectionTriggerType.MANUAL_ALL,
        status=CollectionJobStatus.PENDING,
        priority=100,
        trigger_context={},
        total_items=0,
        succeeded_items=0,
        failed_items=0,
        skipped_items=0,
        created_at=created,
        updated_at=created,
    )

    projection = CollectionJobService.project(job, include_items=False)

    assert projection["created_at"] == "2026-10-07T15:30:00+00:00"
    assert projection["created_at_display"] == "2026-10-08 00:30"


def test_contact_run_source_review_and_export_use_one_kst_projection(tmp_path, monkeypatch):
    db_path = tmp_path / "projection.sqlite3"
    database_url = f"sqlite:///{db_path.as_posix()}"
    monkeypatch.setenv("PUBLICDB2_DATABASE_URL", database_url)
    command.upgrade(Config("alembic.ini"), "head")
    engine = create_db_engine(database_url)
    factory = create_session_factory(engine)
    instant = datetime(2026, 10, 7, 8, 22, tzinfo=timezone.utc)

    try:
        with factory() as session:
            agency_projection, _ = AgencyService(session).create_agency(
                official_name="KST Agency", agency_type=AgencyType.OTHER
            )
            agency_id = uuid.UUID(agency_projection["id"])
            binding_projection, _ = SourceService(session).register_binding(
                url="https://example.org/kst", agency_id=agency_id
            )
            source_id = uuid.UUID(binding_projection["source_id"])
            binding_id = uuid.UUID(binding_projection["binding_id"])
            run = CrawlRun(
                source_id=source_id, status=RunStatus.SUCCESS,
                connection_status=StageStatus.SUCCESS, raw_status=StageStatus.SUCCESS,
                extraction_status=StageStatus.SUCCESS, started_at=instant,
                finished_at=instant, records_observed=1, http_status=200,
            )
            session.add(run)
            session.flush()
            observation = Observation(
                crawl_run_id=run.id, source_id=source_id, observed_at=instant,
                page_url="https://example.org/kst",
            )
            contact = ContactPoint(
                agency_id=agency_id, contact_type=ContactType.PHONE,
                value="02-0000-0000", normalized_value="0200000000",
                verified_at=instant, active=True,
            )
            session.add_all((observation, contact))
            session.flush()
            session.add_all((
                SourceOccurrence(
                    observation_id=observation.id, entity_type=EntityType.CONTACT_POINT,
                    entity_id=contact.id, observed_at=instant,
                ),
                ContactHistory(
                    contact_id=contact.id, value=contact.value,
                    normalized_value=contact.normalized_value, active=True,
                    valid_from=instant,
                ),
            ))
            detection = ChangeDetection(
                source_id=source_id, observation_id=observation.id,
                agency_id=agency_id, coverage_mode=SourceCoverageMode.UNKNOWN,
                status=RunStatus.SUCCESS, candidates_found=1,
                started_at=instant, finished_at=instant,
            )
            session.add(detection)
            session.flush()
            session.add(DetectedChangeCandidate(
                detection_run_id=detection.id, source_id=source_id,
                observation_id=observation.id, agency_id=agency_id,
                entity_type=EntityType.CONTACT_POINT,
                proposed_event_type=ChangeEventType.CONTACT_CHANGED,
                old_value={"value": "old"}, new_value={"value": "new"},
                reason="fixture", review_status=ReviewStatus.PENDING_REVIEW,
                candidate_key="kst-fixture", actionable=True,
            ))
            session.commit()

            contact_detail = ContactService(session).list_page()["details"][0]
            assert contact_detail["verified"] == "2026-10-07 17:22"
            assert contact_detail["provenance"][0]["verified"] == "2026-10-07 17:22"
            assert contact_detail["history"][0].startswith("2026-10-07 17:22 / ")

            source = SourceService(session).get_binding(binding_id)
            assert source["checked"] == "2026-10-07 17:22"
            run_page = RunService(session).list_page()
            assert run_page["items"][0]["time"] == "2026-10-07 17:22"
            assert run_page["details"][0]["started_at"] == "2026-10-07 17:22"
            assert ReviewReadService(session).list_page()["items"][0]["detected"] == "2026-10-07 17:22"

            exported = ContactExportService(
                session, export_root=tmp_path / "exports"
            ).export()
            workbook = load_workbook(exported, read_only=True)
            assert next(workbook.active.iter_rows(min_row=2, values_only=True))[6] == "2026-10-07 17:22"
            workbook.close()
    finally:
        engine.dispose()
