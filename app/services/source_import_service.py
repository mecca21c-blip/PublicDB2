"""Two-phase XLSX/CSV source import with DB revalidation on confirm."""

from __future__ import annotations

import csv
import hashlib
import io
import os
import threading
import time
import uuid
from dataclasses import dataclass
from datetime import datetime
from enum import Enum
from math import ceil
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.orm import Session

from app.core.config import RuntimePaths, ensure_runtime_directories, runtime_paths
from app.models import (
    Agency, AgencyType, ApiSourceKind, CollectionMethod, CollectionTriggerType,
    OrgUnit, OrgUnitType, Source, SourceBinding, SourceImportLog,
)
from app.models.common import utc_now
from app.core.time_presentation import serialize_utc_datetime
from app.services.agency_service import AgencyService
from app.services.collection_job_service import CollectionJobService, MANUAL_PRIORITY
from app.services.normalization import (
    SourceURLValidationError,
    collapse_whitespace,
    normalize_agency_name,
    normalize_org_unit_name,
    normalize_source_url,
)
from app.services.operation_claim_service import OperationClaimService
from app.services.regions import REGIONS
from app.services.source_service import SourceService


MAX_IMPORT_ROWS = 10_000
MAX_IMPORT_BYTES = 20 * 1024 * 1024
PREVIEW_PAGE_SIZE = 50
MAX_PREVIEW_PAGE_SIZE = 100
ALLOWED_SUFFIXES = {".xlsx", ".csv"}
HEADER_ALIASES = {
    "agency": {"기관명", "기관", "agency", "agencyname", "agency_name"},
    "agency_type": {"기관유형", "기관종류", "agencytype", "agency_type"},
    "region": {"지역", "지역코드", "region", "regioncode", "region_code"},
    "org_unit": {"부서명", "부서", "조직명", "department", "orgunit", "org_unit"},
    "url": {"url", "주소", "sourceurl", "source_url", "수집url"},
    "description": {"소스설명", "설명", "description", "source_description"},
    "collection_method": {"수집방식", "collectionmethod", "collection_method", "방식"},
    "scheduled_refresh": {"자동수집", "예약수집", "scheduledrefresh", "scheduled_refresh", "scheduled_refresh_enabled"},
}
IMPORT_METHODS: dict[str, tuple[CollectionMethod, ApiSourceKind | None]] = {
    "": (CollectionMethod.WEB_PAGE, None),
    "스크래핑": (CollectionMethod.WEB_PAGE, None),
    "개별url·스크래핑": (CollectionMethod.WEB_PAGE, None),
    "개별url스크래핑": (CollectionMethod.WEB_PAGE, None),
    "web_page": (CollectionMethod.WEB_PAGE, None),
    "크롤링": (CollectionMethod.WEB_CRAWL, None),
    "indexurl·크롤링": (CollectionMethod.WEB_CRAWL, None),
    "indexurl크롤링": (CollectionMethod.WEB_CRAWL, None),
    "web_crawl": (CollectionMethod.WEB_CRAWL, None),
    "rss": (CollectionMethod.API, ApiSourceKind.RSS),
    "atom": (CollectionMethod.API, ApiSourceKind.ATOM),
}
OPENAPI_METHODS = {"open_api", "openapi", "공개api", "공개_api"}
AGENCY_TYPE_ALIASES = {
    **{item.value.casefold(): item for item in AgencyType},
    "중앙행정기관": AgencyType.CENTRAL_GOVERNMENT,
    "중앙부처": AgencyType.CENTRAL_GOVERNMENT,
    "청": AgencyType.AGENCY,
    "처": AgencyType.AGENCY,
    "행정기관": AgencyType.AGENCY,
    "위원회": AgencyType.COMMISSION,
    "광역자치단체": AgencyType.METROPOLITAN_GOVERNMENT,
    "광역지자체": AgencyType.METROPOLITAN_GOVERNMENT,
    "기초자치단체": AgencyType.BASIC_LOCAL_GOVERNMENT,
    "기초지자체": AgencyType.BASIC_LOCAL_GOVERNMENT,
    "공공기관": AgencyType.PUBLIC_INSTITUTION,
    "기타": AgencyType.OTHER,
}
REGION_ALIASES = {
    **{code.casefold(): code for code, _label in REGIONS},
    **{label.casefold(): code for code, label in REGIONS},
    "서울": "SEOUL", "부산": "BUSAN", "대구": "DAEGU", "인천": "INCHEON",
    "광주": "GWANGJU", "대전": "DAEJEON", "울산": "ULSAN", "세종": "SEJONG",
    "경기": "GYEONGGI", "강원": "GANGWON", "충북": "CHUNGBUK", "충남": "CHUNGNAM",
    "전북": "JEONBUK", "전남": "JEONNAM", "경북": "GYEONGBUK", "경남": "GYEONGNAM",
    "제주": "JEJU",
}
TRUE_ALIASES = {"y", "yes", "true", "1", "포함", "사용", "자동"}
FALSE_ALIASES = {"n", "no", "false", "0", "제외", "미사용"}
LOCAL_AGENCY_TYPES = {AgencyType.METROPOLITAN_GOVERNMENT, AgencyType.BASIC_LOCAL_GOVERNMENT}


class ImportCategory(str, Enum):
    READY = "READY"
    NEW_AGENCY = "NEW_AGENCY"
    NEW_ORG_UNIT = "NEW_ORG_UNIT"
    EXACT_DUPLICATE = "EXACT_DUPLICATE"
    EXISTING_SOURCE_NEW_BINDING = "EXISTING_SOURCE_NEW_BINDING"
    CONFLICT = "CONFLICT"
    INVALID = "INVALID"


class SourceImportError(ValueError):
    pass


@dataclass(frozen=True)
class PreviewFile:
    token: str
    original_filename: str
    temp_path: Path
    sha256: str


class PreviewStore:
    def __init__(self, paths: RuntimePaths | None = None) -> None:
        self.paths = paths or runtime_paths()
        self._entries: dict[str, PreviewFile] = {}
        self._lock = threading.Lock()
        self.cleanup_expired()

    def save(self, filename: str, content: bytes) -> PreviewFile:
        self.cleanup_expired()
        suffix = Path(filename).suffix.lower()
        if suffix not in ALLOWED_SUFFIXES:
            raise SourceImportError("지원 형식은 .xlsx와 .csv입니다.")
        ensure_runtime_directories(self.paths)
        token = uuid.uuid4().hex
        path = self.paths.temp_root / f"source-import-{token}{suffix}"
        path.write_bytes(content)
        entry = PreviewFile(token, Path(filename).name[:500], path, hashlib.sha256(content).hexdigest())
        with self._lock:
            self._entries[token] = entry
        return entry

    def get(self, token: str) -> PreviewFile:
        with self._lock:
            entry = self._entries.get(token)
        if entry is None or not entry.temp_path.is_file():
            raise SourceImportError("미리보기가 없거나 만료되었습니다. 파일을 다시 선택하세요.")
        return entry

    def discard(self, token: str) -> None:
        with self._lock:
            entry = self._entries.pop(token, None)
        if entry is not None:
            entry.temp_path.unlink(missing_ok=True)

    def cleanup_expired(self, max_age_seconds: int = 3600) -> None:
        cutoff = time.time() - max_age_seconds
        with self._lock:
            expired = [token for token, entry in self._entries.items() if not entry.temp_path.exists() or entry.temp_path.stat().st_mtime < cutoff]
        for token in expired:
            self.discard(token)
        if self.paths.temp_root.exists():
            for path in self.paths.temp_root.glob('source-import-*'):
                try:
                    if path.is_file() and path.stat().st_mtime < cutoff:
                        path.unlink(missing_ok=True)
                except OSError:
                    continue


def _header_key(value: Any) -> str:
    return collapse_whitespace(str(value or "")).lower().replace(" ", "")


def _column_map(headers: list[Any]) -> dict[str, int]:
    mapped: dict[str, int] = {}
    for index, header in enumerate(headers):
        key = _header_key(header)
        for field, aliases in HEADER_ALIASES.items():
            if key in aliases and field not in mapped:
                mapped[field] = index
    missing = [name for name in ("agency", "url") if name not in mapped]
    if missing:
        labels = {"agency": "기관명", "url": "URL"}
        raise SourceImportError("필수 컬럼이 없습니다: " + ", ".join(labels[name] for name in missing))
    return mapped


def _cell(values: list[Any], index: int | None) -> str:
    if index is None or index >= len(values) or values[index] is None:
        return ""
    return collapse_whitespace(str(values[index]))


def _compact(value: str) -> str:
    return collapse_whitespace(value).replace(" ", "").casefold()


def _agency_type(value: str) -> AgencyType | None:
    if not value:
        return None
    selected = AGENCY_TYPE_ALIASES.get(_compact(value))
    if selected is None:
        raise SourceImportError("지원되는 기관 유형을 입력하세요.")
    return selected


def _region_code(value: str) -> str | None:
    if not value:
        return None
    selected = REGION_ALIASES.get(_compact(value))
    if selected is None:
        raise SourceImportError("지원되는 지역명 또는 지역 코드를 입력하세요.")
    return selected


def _scheduled_refresh(value: str) -> bool:
    key = _compact(value)
    if not key or key in TRUE_ALIASES:
        return True
    if key in FALSE_ALIASES:
        return False
    raise SourceImportError("자동수집 값은 Y/N, TRUE/FALSE, 포함/제외 중 하나여야 합니다.")


def _method(value: str) -> tuple[CollectionMethod, ApiSourceKind | None]:
    key = _compact(value)
    if key in OPENAPI_METHODS:
        raise SourceImportError(
            "일반 OpenAPI는 전용 Source 설정에서 등록하세요. RSS/Atom은 대량 업로드를 지원합니다."
        )
    selected = IMPORT_METHODS.get(key)
    if selected is None:
        raise SourceImportError("수집 방식은 WEB_PAGE, WEB_CRAWL, RSS, ATOM 중 하나여야 합니다.")
    return selected


def _method_signature(method: CollectionMethod, kind: ApiSourceKind | None) -> tuple[str, str | None]:
    return method.value, kind.value if kind else None


def _source_signature(source: Source) -> tuple[str, str | None]:
    kind = source.api_config.kind if source.collection_method is CollectionMethod.API and source.api_config else None
    return _method_signature(source.collection_method, kind)


def parse_import_file(filename: str, content: bytes, row_limit: int = MAX_IMPORT_ROWS) -> list[dict[str, Any]]:
    suffix = Path(filename).suffix.lower()
    if suffix not in ALLOWED_SUFFIXES:
        raise SourceImportError("지원 형식은 .xlsx와 .csv입니다.")
    if not content:
        raise SourceImportError("빈 파일은 가져올 수 없습니다.")
    if len(content) > MAX_IMPORT_BYTES:
        raise SourceImportError("파일 크기는 20MB를 초과할 수 없습니다.")
    if suffix == ".xlsx":
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheet = workbook.active
            iterator = sheet.iter_rows(values_only=True)
            headers = list(next(iterator))
            source_rows = ((number, list(values)) for number, values in enumerate(iterator, start=2))
        except (StopIteration, OSError, ValueError, KeyError) as error:
            raise SourceImportError("XLSX 파일을 읽을 수 없습니다.") from error
    else:
        try:
            try:
                text = content.decode("utf-8-sig")
            except UnicodeDecodeError:
                text = content.decode("cp949")
            iterator = csv.reader(io.StringIO(text))
            headers = next(iterator)
            source_rows = ((number, list(values)) for number, values in enumerate(iterator, start=2))
        except (StopIteration, UnicodeDecodeError, csv.Error) as error:
            raise SourceImportError("CSV 파일을 읽을 수 없습니다.") from error
    columns = _column_map(headers)
    rows: list[dict[str, Any]] = []
    for row_number, values in source_rows:
        if not any(collapse_whitespace(str(value or "")) for value in values):
            continue
        if len(rows) >= row_limit:
            raise SourceImportError(f"가져오기 행은 최대 {row_limit:,}개입니다.")
        rows.append({
            "row_number": row_number,
            "raw_agency": _cell(values, columns["agency"]),
            "raw_agency_type": _cell(values, columns.get("agency_type")),
            "raw_region": _cell(values, columns.get("region")),
            "raw_org_unit": _cell(values, columns.get("org_unit")),
            "raw_url": _cell(values, columns["url"]),
            "description": _cell(values, columns.get("description")),
            "raw_collection_method": _cell(values, columns.get("collection_method")),
            "raw_scheduled_refresh": _cell(values, columns.get("scheduled_refresh")),
        })
    if not rows:
        raise SourceImportError("가져올 데이터 행이 없습니다.")
    return rows


class SourceImportService:
    def __init__(self, session: Session, paths: RuntimePaths | None = None) -> None:
        self.session = session
        self.paths = paths or runtime_paths()

    def preview(
        self, filename: str, content: bytes, *, filter_name: str = "all",
        page: int = 1, page_size: int = PREVIEW_PAGE_SIZE,
    ) -> dict[str, Any]:
        classified = self.classify(parse_import_file(filename, content))
        return self.preview_page(classified, filter_name=filter_name, page=page, page_size=page_size)

    @staticmethod
    def preview_page(
        classified: dict[str, Any], *, filter_name: str = "all",
        page: int = 1, page_size: int = PREVIEW_PAGE_SIZE,
    ) -> dict[str, Any]:
        predicates = {
            "all": lambda _row: True,
            "importable": lambda row: row["importable"],
            "new_agency": lambda row: row["classification"] == ImportCategory.NEW_AGENCY.value,
            "new_org_unit": lambda row: row["classification"] == ImportCategory.NEW_ORG_UNIT.value,
            "existing_source": lambda row: row["classification"] == ImportCategory.EXISTING_SOURCE_NEW_BINDING.value,
            "duplicate": lambda row: row["classification"] == ImportCategory.EXACT_DUPLICATE.value,
            "conflict": lambda row: row["classification"] == ImportCategory.CONFLICT.value,
            "invalid": lambda row: row["classification"] == ImportCategory.INVALID.value,
        }
        if filter_name not in predicates:
            raise SourceImportError("지원되는 미리보기 필터를 선택하세요.")
        safe_size = min(max(int(page_size or PREVIEW_PAGE_SIZE), 1), MAX_PREVIEW_PAGE_SIZE)
        selected = [row for row in classified["rows"] if predicates[filter_name](row)]
        total = len(selected)
        pages = max(1, ceil(total / safe_size))
        safe_page = min(max(int(page or 1), 1), pages)
        offset = (safe_page - 1) * safe_size
        return {
            "rows": selected[offset:offset + safe_size],
            "summary": classified["summary"],
            "preview": {
                "filter": filter_name, "page": safe_page, "page_size": safe_size,
                "total": total, "pages": pages,
                "has_previous": safe_page > 1, "has_next": safe_page < pages,
            },
        }

    def classify(self, rows: list[dict[str, Any]]) -> dict[str, Any]:
        agencies = list(self.session.scalars(select(Agency).where(Agency.active.is_(True))))
        units = list(self.session.scalars(select(OrgUnit).where(OrgUnit.active.is_(True))))
        sources = list(self.session.scalars(select(Source)))
        bindings = list(self.session.scalars(select(SourceBinding)))
        agency_map: dict[str, list[Agency]] = {}
        for agency in agencies:
            agency_map.setdefault(agency.normalized_name, []).append(agency)
        unit_map: dict[tuple[uuid.UUID, str], list[OrgUnit]] = {}
        for unit in units:
            unit_map.setdefault((unit.agency_id, unit.normalized_name), []).append(unit)
        source_map = {source.normalized_url: source for source in sources}
        binding_keys = {(binding.source_id, binding.scope_key) for binding in bindings}
        seen_signatures: dict[str, tuple[str, str | None]] = {}
        seen_bindings: set[tuple[str, str, str, tuple[str, str | None]]] = set()
        planned_agencies: dict[str, tuple[AgencyType, str | None]] = {}
        planned_units: set[tuple[str, str]] = set()
        projected: list[dict[str, Any]] = []

        for raw in rows:
            item = dict(raw)
            item.update({
                "normalized_url": None, "resolved_agency": None, "resolved_org_unit": None,
                "agency_id": None, "org_unit_id": None, "agency_type": None,
                "region_code": None, "collection_method": None, "api_kind": None,
                "method_config": {}, "scheduled_refresh_enabled": True,
                "will_create_agency": False, "will_create_org_unit": False,
            })
            category, message, importable = self._classify_row(
                item, agency_map, unit_map, source_map, binding_keys,
                seen_signatures, seen_bindings, planned_agencies, planned_units,
            )
            item.update({"classification": category.value, "message": message, "importable": importable})
            projected.append(item)

        category_counts = {category.value: sum(item["classification"] == category.value for item in projected) for category in ImportCategory}
        new_source_urls = {
            item["normalized_url"] for item in projected
            if item["importable"] and item["normalized_url"] not in source_map
        }
        summary = {
            "total": len(projected),
            "importable": sum(item["importable"] for item in projected),
            "new_agencies": sum(item["will_create_agency"] for item in projected if item["importable"]),
            "new_org_units": sum(item["will_create_org_unit"] for item in projected if item["importable"]),
            "new_sources": len(new_source_urls),
            "existing_source_new_bindings": category_counts[ImportCategory.EXISTING_SOURCE_NEW_BINDING.value],
            "duplicates": category_counts[ImportCategory.EXACT_DUPLICATE.value],
            "conflicts": category_counts[ImportCategory.CONFLICT.value],
            "invalid": category_counts[ImportCategory.INVALID.value],
            "categories": category_counts,
        }
        return {"rows": projected, "summary": summary}

    def _classify_row(
        self, item: dict[str, Any], agency_map: dict[str, list[Agency]],
        unit_map: dict[tuple[uuid.UUID, str], list[OrgUnit]], source_map: dict[str, Source],
        binding_keys: set[tuple[uuid.UUID, str]],
        seen_signatures: dict[str, tuple[str, str | None]],
        seen_bindings: set[tuple[str, str, str, tuple[str, str | None]]],
        planned_agencies: dict[str, tuple[AgencyType, str | None]],
        planned_units: set[tuple[str, str]],
    ) -> tuple[ImportCategory, str, bool]:
        agency_name = normalize_agency_name(item["raw_agency"])
        if not agency_name:
            return ImportCategory.INVALID, "기관명은 필수입니다.", False
        if not item["raw_url"]:
            return ImportCategory.INVALID, "URL은 필수입니다.", False
        try:
            normalized_url = normalize_source_url(item["raw_url"])
            imported_type = _agency_type(item.get("raw_agency_type", ""))
            imported_region = _region_code(item.get("raw_region", ""))
            scheduled = _scheduled_refresh(item.get("raw_scheduled_refresh", ""))
            method, api_kind = _method(item.get("raw_collection_method", ""))
        except (SourceURLValidationError, SourceImportError) as error:
            return ImportCategory.INVALID, str(error), False
        item.update({
            "normalized_url": normalized_url,
            "collection_method": method.value,
            "api_kind": api_kind.value if api_kind else None,
            "method_config": {"kind": api_kind.value} if api_kind else {},
            "scheduled_refresh_enabled": scheduled,
        })
        signature = _method_signature(method, api_kind)
        previous_signature = seen_signatures.get(normalized_url)
        if previous_signature is not None and previous_signature != signature:
            return ImportCategory.CONFLICT, "파일 내부에서 같은 canonical URL의 수집 방식 또는 API kind가 충돌합니다.", False
        source = source_map.get(normalized_url)
        if source is not None and _source_signature(source) != signature:
            return ImportCategory.CONFLICT, "기존 canonical URL의 수집 방식 또는 API kind와 충돌합니다.", False

        agency_matches = agency_map.get(agency_name, [])
        if len(agency_matches) > 1:
            return ImportCategory.CONFLICT, "정규화 기관명이 여러 기존 기관과 충돌합니다.", False
        agency = agency_matches[0] if agency_matches else None
        if agency is not None:
            if imported_type is not None and imported_type is not agency.agency_type:
                return ImportCategory.CONFLICT, "가져온 기관 유형이 기존 기관 정보와 충돌합니다.", False
            if imported_region is not None and imported_region != agency.region_code:
                return ImportCategory.CONFLICT, "가져온 지역이 기존 기관 정보와 충돌합니다.", False
            resolved_type, resolved_region = agency.agency_type, agency.region_code
        else:
            desired_type = imported_type or AgencyType.OTHER
            desired_region = imported_region
            planned = planned_agencies.get(agency_name)
            if planned is not None:
                if imported_type is not None and imported_type is not planned[0]:
                    return ImportCategory.CONFLICT, "파일 내부에서 같은 신규 기관의 기관 유형이 충돌합니다.", False
                if imported_region is not None and imported_region != planned[1]:
                    return ImportCategory.CONFLICT, "파일 내부에서 같은 신규 기관의 지역이 충돌합니다.", False
                desired_type, desired_region = planned
            elif desired_type in LOCAL_AGENCY_TYPES and desired_region is None:
                return ImportCategory.INVALID, "신규 광역/기초 자치단체는 지역이 필요합니다.", False
            else:
                planned_agencies[agency_name] = (desired_type, desired_region)
                item["will_create_agency"] = True
            resolved_type, resolved_region = desired_type, desired_region
        item.update({
            "resolved_agency": agency.official_name if agency else collapse_whitespace(item["raw_agency"]),
            "agency_id": str(agency.id) if agency else None,
            "agency_type": resolved_type.value,
            "region_code": resolved_region,
        })

        org_name = normalize_org_unit_name(item["raw_org_unit"])
        unit = None
        agency_identity = f"id:{agency.id}" if agency else f"new:{agency_name}"
        if org_name and agency:
            matches = unit_map.get((agency.id, org_name), [])
            if len(matches) > 1:
                return ImportCategory.CONFLICT, "기관 내 정규화 부서명이 여러 부서와 충돌합니다.", False
            unit = matches[0] if matches else None
        if org_name and unit is None:
            planned_unit_key = (agency_identity, org_name)
            if planned_unit_key not in planned_units:
                planned_units.add(planned_unit_key)
                item["will_create_org_unit"] = True
        item["resolved_org_unit"] = unit.name if unit else (collapse_whitespace(item["raw_org_unit"]) or None)
        item["org_unit_id"] = str(unit.id) if unit else None

        org_identity = f"id:{unit.id}" if unit else (f"new:{org_name}" if org_name else "agency")
        semantic_key = (normalized_url, agency_identity, org_identity, signature)
        if semantic_key in seen_bindings:
            item["will_create_agency"] = False
            item["will_create_org_unit"] = False
            return ImportCategory.EXACT_DUPLICATE, "파일 내부에서 같은 URL과 기관/부서가 중복되었습니다.", False
        seen_signatures[normalized_url] = signature
        seen_bindings.add(semantic_key)

        if source and agency and (
            source.id, f"org:{unit.id}" if unit else f"agency:{agency.id}"
        ) in binding_keys and not item["will_create_org_unit"]:
            item["will_create_agency"] = False
            return ImportCategory.EXACT_DUPLICATE, "같은 URL과 기관/부서 연결이 이미 등록되어 있습니다.", False
        if agency is None:
            return ImportCategory.NEW_AGENCY, "기관을 새로 만들고 URL을 연결합니다.", True
        if org_name and unit is None:
            return ImportCategory.NEW_ORG_UNIT, "기관 아래 부서를 새로 만들고 URL을 연결합니다.", True
        if source:
            return ImportCategory.EXISTING_SOURCE_NEW_BINDING, "기존 canonical URL에 새 기관/부서 연결을 만듭니다.", True
        return ImportCategory.READY, "새 URL과 연결을 등록합니다.", True

    def confirm(
        self, entry: PreviewFile, *, collect_after_confirm: bool = False,
        requested_by_user_id: uuid.UUID | None = None,
    ) -> dict[str, Any]:
        content = entry.temp_path.read_bytes()
        classified = self.classify(parse_import_file(entry.original_filename, content))
        preexisting_urls = set(self.session.scalars(select(Source.normalized_url)))
        counts = {
            "input_rows": classified["summary"]["total"], "created_agencies": 0,
            "created_org_units": 0, "created_sources": 0, "created_bindings": 0,
            "existing_source_new_bindings": 0, "duplicates_skipped": classified["summary"]["duplicates"],
            "invalid_rows": classified["summary"]["invalid"], "conflicts": classified["summary"]["conflicts"],
            "unexpected_failures": 0,
        }
        self.session.rollback()
        claims = OperationClaimService(self.session)
        acquired = claims.acquire(f"import:{entry.token}", "SOURCE_IMPORT_CONFIRM")
        if not acquired.acquired:
            if acquired.claim.status == "COMPLETED" and acquired.claim.result_payload:
                return {**acquired.claim.result_payload, "reused": True}
            raise SourceImportError("이 가져오기 미리보기는 이미 확인 처리 중입니다.")
        retained_path: Path | None = None
        registration_committed = False
        created_source_ids: set[uuid.UUID] = set()
        try:
            retained_path = self._retain_file(entry, content)
            agency_service = AgencyService(self.session)
            source_service = SourceService(self.session)
            for row in classified["rows"]:
                if not row["importable"]:
                    continue
                try:
                    with self.session.begin_nested():
                        normalized_agency = normalize_agency_name(row["resolved_agency"])
                        agency = self.session.scalar(select(Agency).where(
                            Agency.normalized_name == normalized_agency, Agency.active.is_(True)
                        ))
                        created_agency = False
                        if agency is None:
                            agency_data, created_agency = agency_service.create_agency(
                                official_name=row["resolved_agency"],
                                agency_type=AgencyType(row["agency_type"]),
                                region_code=row["region_code"], commit=False,
                            )
                            agency = self.session.get(Agency, uuid.UUID(agency_data["id"]))
                        agency_id = agency.id

                        org_unit_id = None
                        created_org_unit = False
                        if row["resolved_org_unit"]:
                            normalized_unit = normalize_org_unit_name(row["resolved_org_unit"])
                            unit = self.session.scalar(select(OrgUnit).where(
                                OrgUnit.agency_id == agency_id,
                                OrgUnit.normalized_name == normalized_unit,
                                OrgUnit.active.is_(True),
                            ))
                            if unit is None:
                                unit_data = agency_service.create_org_unit(
                                    agency_id=agency_id, name=row["resolved_org_unit"],
                                    unit_type=OrgUnitType.DEPARTMENT, commit=False,
                                )
                                unit = self.session.get(OrgUnit, uuid.UUID(unit_data["id"]))
                                created_org_unit = True
                            org_unit_id = unit.id

                        existing_source_id = self.session.scalar(select(Source.id).where(
                            Source.normalized_url == row["normalized_url"]
                        ))
                        binding, created_binding = source_service.register_binding(
                            url=row["raw_url"], agency_id=agency_id, org_unit_id=org_unit_id,
                            description=row["description"], collection_method=row["collection_method"],
                            method_config=row["method_config"],
                            scheduled_refresh_enabled=row["scheduled_refresh_enabled"], commit=False,
                        )
                        source_id = uuid.UUID(binding["source_id"])
                        if existing_source_id is None:
                            created_source_ids.add(source_id)
                        counts["created_agencies"] += int(created_agency)
                        counts["created_org_units"] += int(created_org_unit)
                        counts["created_sources"] += int(existing_source_id is None)
                        counts["created_bindings"] += int(created_binding)
                        counts["existing_source_new_bindings"] += int(
                            row["normalized_url"] in preexisting_urls and created_binding
                        )
                except Exception:
                    counts["unexpected_failures"] += 1

            confirmed_at = utc_now()
            relative = retained_path.relative_to(self.paths.project_root).as_posix()
            import_log = SourceImportLog(
                original_filename=entry.original_filename, stored_path=relative,
                sha256=entry.sha256, confirmed_at=confirmed_at,
                input_rows=counts["input_rows"], result_counts=dict(counts),
            )
            self.session.add(import_log)
            self.session.commit()
            registration_committed = True

            collection_job_id = None
            collection_source_count = len(created_source_ids)
            collection_queue_status = "NOT_REQUESTED"
            if collect_after_confirm:
                if created_source_ids:
                    try:
                        sources = list(self.session.scalars(select(Source).where(Source.id.in_(created_source_ids))))
                        job = CollectionJobService(self.session).create_from_sources(
                            CollectionTriggerType.MANUAL_SELECTION, sources,
                            priority=MANUAL_PRIORITY, requested_by_user_id=requested_by_user_id,
                            trigger_context={
                                "scope": CollectionTriggerType.MANUAL_SELECTION.value,
                                "origin": "SOURCE_IMPORT",
                                "import_filename": entry.original_filename,
                                "import_sha256": entry.sha256,
                                "import_log_id": str(import_log.id),
                                "new_source_count": collection_source_count,
                            },
                        )
                        collection_job_id = str(job.id)
                        collection_queue_status = "QUEUED"
                    except Exception:
                        self.session.rollback()
                        collection_queue_status = "FAILED_TO_QUEUE"
                else:
                    collection_queue_status = "NO_NEW_SOURCES"

            durable_result = {
                "summary": counts,
                "file": {
                    "original_filename": entry.original_filename, "stored_path": relative,
                    "sha256": entry.sha256, "confirmed_at": serialize_utc_datetime(confirmed_at),
                },
                "collection_requested": collect_after_confirm,
                "collection_job_id": collection_job_id,
                "collection_source_count": collection_source_count,
                "collection_queue_status": collection_queue_status,
            }
            import_log = self.session.get(SourceImportLog, import_log.id)
            import_log.result_counts = {
                **counts, "collection_requested": collect_after_confirm,
                "collection_job_id": collection_job_id,
                "collection_source_count": collection_source_count,
                "collection_queue_status": collection_queue_status,
            }
            claim = self.session.get(type(acquired.claim), acquired.claim.id)
            claims.complete(claim, durable_result)
            self.session.commit()
        except Exception:
            self.session.rollback()
            if retained_path is not None and not registration_committed:
                retained_path.unlink(missing_ok=True)
            raise
        return {**durable_result, "reused": False}

    def _retain_file(self, entry: PreviewFile, content: bytes) -> Path:
        now = datetime.now().astimezone()
        destination = self.paths.import_root / f"{now:%Y}" / f"{now:%m}" / f"{now:%d}"
        destination.mkdir(parents=True, exist_ok=True)
        suffix = Path(entry.original_filename).suffix.lower()
        target = destination / f"{now:%Y%m%dT%H%M%S}-{uuid.uuid4().hex}{suffix}"
        temporary = target.with_suffix(target.suffix + '.tmp')
        try:
            with temporary.open('xb') as stream:
                stream.write(content)
                stream.flush()
                os.fsync(stream.fileno())
            os.replace(temporary, target)
        finally:
            temporary.unlink(missing_ok=True)
        return target

