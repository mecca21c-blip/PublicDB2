from __future__ import annotations

import io
import uuid

from fastapi import APIRouter, Depends, File, HTTPException, Request, UploadFile, status
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from sqlalchemy.orm import Session

from app.api.dependencies import get_session, require_operator, require_viewer
from app.api.schemas import (
    ExclusionRequest, InteractiveSourceBatchRequest, ScrapeBatchRequest,
    SourceBindingCreate, SourceBindingUpdate,
)
from app.models import ApiAuthMode, CollectionMethod, User, UserRole
from app.services.normalization import SourceURLValidationError
from app.services.source_service import SourceBindingConflict, SourceService, SourceServiceError
from app.services.source_method_service import MethodConfigError
from app.services.source_import_service import (
    MAX_IMPORT_BYTES, PREVIEW_PAGE_SIZE, SourceImportError, SourceImportService,
)
from app.services.source_query_service import FILTER_METHODS, SourceFilterError, SourceFilterSpec, SourceQueryService
from app.services.scrape_batch_service import ScrapeBatchError, ScrapeBatchService
from app.services.source_interactive_batch_service import InteractiveBatchError, SourceInteractiveBatchService


router = APIRouter(prefix="/api/source-bindings", tags=["sources"], dependencies=[Depends(require_viewer)])


def _failure(error: ValueError) -> HTTPException:
    code = status.HTTP_409_CONFLICT if isinstance(error, SourceBindingConflict) else status.HTTP_400_BAD_REQUEST
    return HTTPException(status_code=code, detail=str(error))


@router.get("")
def list_bindings(
    search: str | None = None,
    region_code: list[str] | None = None,
    agency_id: uuid.UUID | None = None,
    org_unit_id: uuid.UUID | None = None,
    methods: str | None = None,
    source_status: str | None = None,
    scheduled: str = "all",
    page: int = 1,
    page_size: int = 100,
    session: Session = Depends(get_session),
) -> dict:
    try:
        parsed_methods = list(FILTER_METHODS) if methods is None else [
            CollectionMethod(value) for value in methods.split(",") if value
        ]
        spec = SourceFilterSpec.build(
            search=search, region_codes=region_code, agency_id=agency_id,
            org_unit_id=org_unit_id, methods=parsed_methods,
            status=source_status, scheduled=scheduled,
        )
        return SourceQueryService(session).list_page(spec, page=page, page_size=page_size)
    except (SourceFilterError, ValueError) as error:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(error)) from error


def _require_secret_config_admin(method, config, user: User) -> None:
    if method is CollectionMethod.API:
        auth = ApiAuthMode((config or {}).get("auth_mode", ApiAuthMode.NONE))
        if auth is not ApiAuthMode.NONE and user.role is not UserRole.ADMIN:
            raise HTTPException(status_code=403, detail="API 자격증명 설정은 관리자만 변경할 수 있습니다.")


@router.post("", status_code=status.HTTP_201_CREATED)
def create_binding(
    payload: SourceBindingCreate, user: User = Depends(require_operator),
    session: Session = Depends(get_session),
) -> dict:
    try:
        _require_secret_config_admin(payload.collection_method, payload.method_config, user)
        item, created = SourceService(session).register_binding(**payload.model_dump())
        return {"item": item, "created": created}
    except (SourceServiceError, SourceURLValidationError, MethodConfigError, ValueError) as error:
        raise _failure(error) from error


@router.post("/scrape-batch/preview", dependencies=[Depends(require_operator)])
def preview_scrape_batch(payload: ScrapeBatchRequest, session: Session = Depends(get_session)) -> dict:
    try:
        return ScrapeBatchService(session).preview(
            urls=payload.urls, agency_id=payload.agency_id,
            org_unit_id=payload.org_unit_id, binding_scope=payload.binding_scope,
        )
    except (ScrapeBatchError, SourceServiceError, SourceURLValidationError, ValueError) as error:
        raise _failure(error) from error


@router.post("/scrape-batch/confirm", status_code=status.HTTP_201_CREATED, dependencies=[Depends(require_operator)])
def confirm_scrape_batch(payload: ScrapeBatchRequest, session: Session = Depends(get_session)) -> dict:
    try:
        return ScrapeBatchService(session).confirm(**payload.model_dump())
    except (ScrapeBatchError, SourceServiceError, SourceURLValidationError, MethodConfigError, ValueError) as error:
        raise _failure(error) from error


@router.post("/interactive/preview")
def preview_interactive_batch(
    payload: InteractiveSourceBatchRequest,
    user: User = Depends(require_operator),
    session: Session = Depends(get_session),
) -> dict:
    try:
        _require_secret_config_admin(payload.collection_method, payload.method_config, user)
        SourceInteractiveBatchService.validate_method_config(
            payload.collection_method, payload.method_config, payload.rows[0].url,
        )
        return SourceInteractiveBatchService(session).preview(payload)
    except (InteractiveBatchError, SourceServiceError, SourceURLValidationError, MethodConfigError, ValueError) as error:
        raise _failure(error) from error


@router.post("/interactive/register", status_code=status.HTTP_201_CREATED)
def register_interactive_batch(
    payload: InteractiveSourceBatchRequest,
    user: User = Depends(require_operator),
    session: Session = Depends(get_session),
) -> dict:
    try:
        _require_secret_config_admin(payload.collection_method, payload.method_config, user)
        return SourceInteractiveBatchService(session).register(payload)
    except (InteractiveBatchError, SourceServiceError, SourceURLValidationError, MethodConfigError, ValueError) as error:
        raise _failure(error) from error


@router.get("/{binding_id}")
def get_binding(binding_id: uuid.UUID, session: Session = Depends(get_session)) -> dict:
    try:
        return {"item": SourceService(session).get_binding(binding_id)}
    except SourceServiceError as error:
        raise _failure(error) from error


@router.patch("/{binding_id}")
def update_binding(
    binding_id: uuid.UUID, payload: SourceBindingUpdate,
    user: User = Depends(require_operator), session: Session = Depends(get_session),
) -> dict:
    try:
        values = payload.model_dump(exclude_unset=True)
        if payload.collection_method is not None:
            _require_secret_config_admin(payload.collection_method, payload.method_config, user)
        return {"item": SourceService(session).update_binding(binding_id, **values)}
    except (SourceServiceError, SourceURLValidationError, MethodConfigError, ValueError) as error:
        raise _failure(error) from error


@router.post("/{binding_id}/exclude", dependencies=[Depends(require_operator)])
def exclude_binding(binding_id: uuid.UUID, payload: ExclusionRequest, session: Session = Depends(get_session)) -> dict:
    try:
        return {"item": SourceService(session).exclude_binding(binding_id, payload.reason)}
    except SourceServiceError as error:
        raise _failure(error) from error


@router.post("/{binding_id}/reactivate", dependencies=[Depends(require_operator)])
def reactivate_binding(binding_id: uuid.UUID, session: Session = Depends(get_session)) -> dict:
    try:
        return {"item": SourceService(session).reactivate_binding(binding_id)}
    except SourceServiceError as error:
        raise _failure(error) from error


@router.post("/imports/preview", dependencies=[Depends(require_operator)])
async def preview_import(request: Request, file: UploadFile = File(...), session: Session = Depends(get_session)) -> dict:
    content = await file.read(MAX_IMPORT_BYTES + 1)
    try:
        service = SourceImportService(session, request.app.state.runtime_paths)
        preview = service.preview(file.filename or "upload", content)
        entry = request.app.state.source_import_previews.save(file.filename or "upload", content)
        return {"token": entry.token, "filename": entry.original_filename, **preview}
    except SourceImportError as error:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(error)) from error


@router.get("/imports/{token}/preview", dependencies=[Depends(require_operator)])
def page_import_preview(
    token: str, request: Request, filter: str = "all", page: int = 1,
    page_size: int = PREVIEW_PAGE_SIZE, session: Session = Depends(get_session),
) -> dict:
    try:
        entry = request.app.state.source_import_previews.get(token)
        return SourceImportService(session, request.app.state.runtime_paths).preview(
            entry.original_filename, entry.temp_path.read_bytes(),
            filter_name=filter, page=page, page_size=page_size,
        )
    except SourceImportError as error:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(error)) from error


@router.post("/imports/{token}/confirm")
def confirm_import(
    token: str, request: Request, collect_after_confirm: bool = False,
    user: User = Depends(require_operator), session: Session = Depends(get_session),
) -> dict:
    store = request.app.state.source_import_previews
    try:
        entry = store.get(token)
        return SourceImportService(session, request.app.state.runtime_paths).confirm(
            entry, collect_after_confirm=collect_after_confirm,
            requested_by_user_id=user.id,
        )
    except SourceImportError as error:
        raise HTTPException(status_code=status.HTTP_400_BAD_REQUEST, detail=str(error)) from error
    finally:
        store.discard(token)


@router.delete("/imports/{token}", dependencies=[Depends(require_operator)])
def discard_import(token: str, request: Request) -> dict[str, bool]:
    request.app.state.source_import_previews.discard(token)
    return {"discarded": True}


@router.get("/imports/template.csv")
def download_import_template() -> Response:
    content = "\ufeff기관명,기관유형,지역,부서명,URL,수집방식,소스설명,자동수집\r\n"
    return Response(content=content.encode("utf-8"), media_type="text/csv; charset=utf-8", headers={"Content-Disposition": 'attachment; filename="publicdb2-source-import-template.csv"'})


@router.get("/imports/template.xlsx")
def download_import_xlsx_template() -> Response:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "수집 소스"
    headers = ["기관명", "기관유형", "지역", "부서명", "URL", "수집방식", "소스설명", "자동수집"]
    sheet.append(headers)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = "A1:H1"
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="235EA7")
    for column, width in zip("ABCDEFGH", (24, 26, 16, 24, 52, 22, 42, 16)):
        sheet.column_dimensions[column].width = width
    guide = workbook.create_sheet("입력 안내")
    guide.append(["항목", "안내"])
    guide.append(["필수", "기관명, URL"])
    guide.append(["기관유형", "CENTRAL_GOVERNMENT, AGENCY, COMMISSION, METROPOLITAN_GOVERNMENT, BASIC_LOCAL_GOVERNMENT, PUBLIC_INSTITUTION, OTHER 또는 한국어 별칭"])
    guide.append(["지역", "PublicDB2 지역 코드 또는 지역명. 광역/기초 자치단체 신규 등록 시 필수"])
    guide.append(["수집방식", "WEB_PAGE, WEB_CRAWL, RSS, ATOM. 일반 OpenAPI는 전용 Source 설정에서 등록"])
    guide.append(["자동수집", "Y/N, TRUE/FALSE, 포함/제외. 빈 값은 포함"])
    guide.column_dimensions["A"].width = 20
    guide.column_dimensions["B"].width = 120
    stream = io.BytesIO()
    workbook.save(stream)
    return Response(
        content=stream.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="publicdb2-source-import-template.xlsx"'},
    )
